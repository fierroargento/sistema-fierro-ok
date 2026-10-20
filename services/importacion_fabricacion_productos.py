"""Adaptador de Planilla General al alta de fabricación existente; nunca ejecuta fórmulas."""
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile, BadZipFile

from itsdangerous import URLSafeTimedSerializer, BadSignature, SignatureExpired

EXCLUIDOS = frozenset(f'L-MDF-{n:03d}' for n in range(1, 6))
MAX_PRODUCTOS = 500
MAX_BYTES = 2 * 1024 * 1024


def texto(v):
    if v is None: return ''
    if isinstance(v, str) and v.startswith('='):
        raise ValueError('Hay una fórmula en un dato de entrada. Usá valores en las columnas de origen.')
    return str(v).strip()


def decimal(v, nombre, *, cero=False):
    try:
        s=texto(v)
        if len(s)>32 or isinstance(v,bool): raise ValueError()
        d=Decimal(s.replace(',', '.'))
        if not d.is_finite() or d<0 or (not cero and d==0) or d>Decimal('999999999999'): raise ValueError()
        return d
    except (InvalidOperation, ValueError):
        raise ValueError(f'{nombre}: completá un número {"positivo o cero" if cero else "mayor que cero"}.') from None


def numero(d, precision=6):
    return format(d.quantize(Decimal(1).scaleb(-precision),rounding=ROUND_HALF_UP),'f')


def leer(archivo):
    nombre=str(getattr(archivo,'filename','')).lower()
    if not nombre.endswith(('.xlsx','.xlsm')): raise ValueError('Elegí la Planilla General en XLSX o XLSM.')
    contenido=archivo.read(MAX_BYTES+1)
    if len(contenido)>MAX_BYTES: raise ValueError('La planilla supera 2 MB.')
    try:
        with ZipFile(BytesIO(contenido)) as z:
            if len(z.infolist())>2000 or sum(i.file_size for i in z.infolist())>30*1024*1024:
                raise ValueError('El libro descomprimido supera el tamaño permitido.')
    except BadZipFile: raise ValueError('El archivo no es un libro Excel válido.') from None
    from openpyxl import load_workbook
    libro=load_workbook(BytesIO(contenido),read_only=True,data_only=False,keep_links=False)
    try:
        requeridas={'Productos':['ID Producto','Nombre','Línea','Descripción','Activo (Sí/No)'],
            'Componentes':['ID Producto','Codigo Proveedor','Nombre Producto','Componente','Descripcion','Medida Cm','Cantidad','Unidad','Total','Cantidad Placa Entera','Tiempo 1/4 placa','Cantidad x Carton'],
            'Tareas':['ID Producto','Tarea','Segundos'],
            'Resumen Producto':['ID Producto']}
        tablas={}
        from services.importacion_productos_costeo import normalizar
        for hoja, columnas in requeridas.items():
            if hoja not in libro.sheetnames: raise ValueError(f'Falta la hoja {hoja}.')
            ws=libro[hoja]
            if ws.max_row>5001 or ws.max_column>60: raise ValueError(f'La hoja {hoja} supera los límites permitidos.')
            fuente=ws.iter_rows(values_only=True)
            encabezados=next(fuente,())
            if [normalizar(x) for x in encabezados[:len(columnas)]] != [normalizar(x) for x in columnas]:
                raise ValueError(f'Revisá los encabezados de la hoja {hoja}.')
            if hoja=='Resumen Producto' and (len(encabezados)<11 or normalizar(encabezados[6])!='margen %' or normalizar(encabezados[10])!='iibb (%)' or normalizar(encabezados[8]) not in ('iva (21%)','iva (%)')):
                raise ValueError('No se reconocen Margen %, IVA e IIBB de Resumen Producto.')
            tablas[hoja]=[{'numero':n,'valores':list(r)} for n,r in enumerate(fuente,2) if r and r[0] not in (None,'')]
        productos=tablas['Productos']
        if not productos or len(productos)>MAX_PRODUCTOS: raise ValueError('El lote debe contener entre 1 y 500 productos.')
        skus={texto(r['valores'][0]).upper() for r in productos}
        for hoja in ('Componentes','Tareas','Resumen Producto'):
            huerfanos={texto(r['valores'][0]).upper() for r in tablas[hoja]}-skus
            if huerfanos: raise ValueError(f'{hoja}: hay SKU sin ficha en Productos: '+', '.join(sorted(huerfanos)[:5]))
        return {'tablas':tablas,'archivo_sha256':sha256(contenido).hexdigest()}
    finally: libro.close()


def _agrupar(tablas, hoja):
    salida=defaultdict(list)
    for r in tablas[hoja]: salida[texto(r['valores'][0]).upper()].append(r)
    return salida


def _opciones(opciones, session, org):
    from models.unidad_negocio import UnidadNegocio
    unidad=str(opciones.get('unidad_id') or '')
    if unidad and (not unidad.isdigit() or session.query(UnidadNegocio).filter_by(id=int(unidad),organizacion_id=org,activa=True).first() is None):
        raise ValueError('Elegí una unidad de fabricación activa de tu organización.')
    return {'unidad_id':int(unidad) if unidad else None,
        **{c:str(opciones.get(c) or '') for c in ('tarifa_herreria','tarifa_laser','maquina_laser')},
        **{f'habilita_{uso}':opciones.get(f'habilita_{uso}') in (True,'1') for uso in ('compra','venta','insumo','componente')}}


def _indice_materiales(session, org, unidad):
    from models.fuentes_costo_productivo import InsumoProductivo
    from services.selector_materiales_producto import enriquecer_insumos
    recursos=session.query(InsumoProductivo).filter_by(organizacion_id=org,activo=True).all()
    base=[{'id':r.id,'codigo':r.codigo,'unidad':r.unidad_medida} for r in recursos if r.unidad_negocio_id in (None,unidad)]
    indice=defaultdict(set)
    for r in enriquecer_insumos(base,session,org):
        for codigo in [r['codigo']]+r['skus']+[c['codigo'] for c in r['codigos_proveedor']]:
            if str(codigo).strip(): indice[str(codigo).strip().casefold()].add(r['id'])
    return indice


def _formulario_adicional(resumen):
    if len(resumen)!=1: raise ValueError('El producto requiere una única fila en Resumen Producto.')
    r=resumen[0]['valores']
    if len(r)<11: raise ValueError('Faltan impuestos o utilidad en Resumen Producto.')
    return dict(datos_generales_presentados='1',costo_maestro_presentado='1',iibb_maestro_presentado='1',
        rubro='',subrubro='',marca='',moneda_costo='ARS',precio_costo='',
        utilidad_base='costo',utilidad_pct=numero(decimal(r[6],'Utilidad',cero=True)*100) if texto(r[6]) else '',
        iva_pct=numero(decimal(r[8],'IVA',cero=True)*100),iibb_pct=numero(decimal(r[10],'IIBB',cero=True)*100),
        proveedores_json='[]',sucursales_json='[]',depositos_json='[]',codigos_barras_json='[]',imagenes_conservar_json='[]',imagen_principal='')


def _receta(componentes,tareas,opciones,indice,linea):
    if not componentes: raise ValueError('Faltan los componentes del producto de fabricación.')
    piezas=[];avisos=[];minutos_laser=Decimal(0)
    for registro in componentes:
        r=registro['valores'];n=registro['numero']
        if len(r)<12: raise ValueError(f'Componentes, fila {n}: faltan columnas.')
        codigo=texto(r[1]);ids=indice.get(codigo.casefold(),set())
        if len(ids)!=1: raise ValueError(f'Componentes, fila {n}: insumo {codigo} no encontrado o ambiguo en la unidad elegida.')
        q=decimal(r[6],f'Cantidad, fila {n}')
        unidad=texto(r[7]).casefold();largo=texto(r[5]);placa=texto(r[9]);carton=texto(r[11])
        p={'codigo_origen':codigo,'nombre':texto(r[4]),'insumo_id':next(iter(ids)),'cantidad':numero(q),'modo':'directo','largo_cm':'','rendimiento':''}
        if placa:
            rendimiento=decimal(placa,f'Rendimiento por placa, fila {n}');p.update(modo='rendimiento',rendimiento=numero(rendimiento))
        elif unidad in ('m','metro','metros'):
            p.update(modo='longitud',largo_cm=numero(decimal(largo,f'Medida en cm, fila {n}')))
        elif not largo and carton:
            p.update(modo='rendimiento',rendimiento=numero(decimal(carton,f'Rendimiento por cartón, fila {n}')))
        elif largo: raise ValueError(f'Componentes, fila {n}: medida sin unidad metro; requiere revisión.')
        piezas.append(p)
        if linea=='laser' and texto(r[10]):
            if not placa: raise ValueError(f'Componentes, fila {n}: el tiempo de 1/4 de placa requiere rendimiento de placa entera.')
            minutos_laser+=decimal(r[10],f'Tiempo láser, fila {n}')*4*q/rendimiento
            if q!=1: avisos.append(f'Fila {n}: tiempo láser proporcional a {numero(q)} piezas; revisar contra la fórmula histórica.')
    operaciones=[]
    for r in tareas:
        vals=r['valores']
        operaciones.append({'nombre':texto(vals[1]),'segundos':numero(decimal(vals[2],f'Tiempo de tarea, fila {r["numero"]}'))})
    maquinas=[]
    tarifa=opciones['tarifa_laser'] if linea=='laser' else opciones['tarifa_herreria']
    if linea=='laser':
        if operaciones: raise ValueError('El producto láser tiene tareas además del tiempo por placa; revisá el posible doble cómputo.')
        if minutos_laser<=0: raise ValueError('Falta el tiempo láser por placa. No se supone mano de obra cero.')
        if not opciones['maquina_laser']: raise ValueError('Seleccioná la máquina láser (costo de uso sin mano de obra).')
        operaciones=[{'nombre':'Corte láser','segundos':numero(minutos_laser*60)}]
        maquinas=[{'maquina_id':opciones['maquina_laser'],'nombre':'Corte láser','minutos':numero(minutos_laser,4),'observacion':'Tiempo por unidad desde rendimiento y tiempo de 1/4 de placa.'}]
    if not operaciones: avisos.append('Sin tareas en el Excel: el costo incluye los insumos cargados (incluida confección tercerizada si corresponde). Revisar antes de fijar precios.')
    if operaciones and not tarifa: raise ValueError('Seleccioná la tarifa de mano de obra correspondiente a la línea.')
    return {'unidad_id':opciones['unidad_id'],'piezas':piezas,'tareas':operaciones,'tarifa_id':tarifa or None,'maquinas':maquinas,'fijos':[]},avisos


def previsualizar(lote, *, opciones, Producto, session, organizacion_id):
    from services.productos_ficha_maestra import normalizar_ficha, validar_vinculos
    from services.productos_datos_generales import leer_formulario_completo, esquema_datos_disponible
    from services.productos_produccion_maestro import preparar, calcular
    from services.productos_despiece import disponible
    if not esquema_datos_disponible(session) or not disponible(session): raise ValueError('Falta preparar las fichas completas de productos en este entorno.')
    opciones=_opciones(opciones,session,organizacion_id)
    tablas=lote['tablas'];componentes=_agrupar(tablas,'Componentes');tareas=_agrupar(tablas,'Tareas');resumen=_agrupar(tablas,'Resumen Producto')
    repetidos=Counter(texto(r['valores'][0]).upper() for r in tablas['Productos'])
    existentes=defaultdict(list)
    for p in session.query(Producto).filter_by(organizacion_id=organizacion_id).all(): existentes[p.sku.strip().upper()].append(p.id)
    indice=_indice_materiales(session,organizacion_id,opciones['unidad_id'])
    vista=[]
    from services.importacion_productos_costeo import normalizar
    from services.productos_ficha_maestra import booleano
    for registro in tablas['Productos']:
        r=registro['valores'];sku=texto(r[0]).upper()
        fila={'numero':registro['numero'],'sku':sku,'nombre':texto(r[1]),'accion':'crear','errores':[],'avisos':[],'datos':None,'adicionales':None,'produccion':None,'costo_centavos':None,'huella_costo':'','piezas':0,'tareas':0,'producto_id':None}
        try:
            if sku in EXCLUIDOS: fila['accion']='excluido'
            elif repetidos[sku]!=1: raise ValueError('SKU repetido en el archivo; no se eligió ninguna fila.')
            elif len(existentes[sku])>1: raise ValueError('El sistema tiene más de un producto con ese SKU; revisá el maestro.')
            elif existentes[sku]: fila.update(accion='conservar',producto_id=existentes[sku][0])
            else:
                linea=normalizar(r[2])
                if linea not in ('herreria','laser'): raise ValueError('La línea requiere un método de fabricación definido (Herrería o Láser).')
                datos=dict(sku=sku,nombre=texto(r[1]),linea=texto(r[2]),descripcion=texto(r[3]),tipo='produccion',unidad_medida='unidad',
                    activo=booleano(r[4],'Activo'),no_stockeable=False,permitir_stock_negativo=False,insumo_id=None,
                    permite_correo=False,permite_via_cargo=False,requiere_revision_logistica=True,observacion_logistica='',
                    **{k:v for k,v in opciones.items() if k.startswith('habilita_')})
                datos=normalizar_ficha(datos);validar_vinculos(None,datos,session,organizacion_id)
                adicionales=_formulario_adicional(resumen[sku]);leer_formulario_completo(adicionales,[])
                receta,avisos=_receta(componentes[sku],tareas[sku],opciones,indice,linea)
                if not adicionales['utilidad_pct']: avisos.append('Utilidad pretendida sin definir en el Excel. Completar antes de calcular precios de venta.')
                plan=preparar(receta,session,organizacion_id);costo=calcular(plan,organizacion_id)
                fila.update(datos=datos,adicionales=adicionales,produccion=receta,avisos=avisos,costo_centavos=costo['total_centavos'],huella_costo=costo['huella'],piezas=len(receta['piezas']),tareas=len(receta['tareas']))
        except (ValueError,IndexError,TypeError) as error:
            fila['accion']='rechazado';fila['errores']=[str(error)]
        vista.append(fila)
    return {'filas':vista,'opciones':opciones,'archivo_sha256':lote['archivo_sha256']}


def firmar(vista, org, usuario, secreto):
    return URLSafeTimedSerializer(secreto,salt='fabricacion-productos-v1').dumps({'org':org,'usuario':usuario,'vista':vista})


def aplicar(token, *, organizacion_id, usuario, secreto, Producto, session):
    if len(token)>500000: raise ValueError('La vista previa supera el tamaño permitido.')
    try: lote=URLSafeTimedSerializer(secreto,salt='fabricacion-productos-v1').loads(token,max_age=900)
    except (BadSignature,SignatureExpired): raise ValueError('La vista previa venció o fue modificada. Volvé a analizar el archivo.') from None
    if lote['org']!=organizacion_id or lote['usuario']!=usuario.id: raise ValueError('El lote no pertenece al usuario y organización activos.')
    vista=lote['vista'];filas=vista['filas']
    if not filas or len(filas)>MAX_PRODUCTOS or any(f['errores'] or f['accion']=='rechazado' for f in filas): raise ValueError('Corregí las filas con errores antes de confirmar. No se aplicó el lote.')
    from services.productos_ficha_maestra import guardar_ficha
    from services.productos_datos_generales import leer_formulario_completo, guardar_datos_generales
    from services.productos_produccion_maestro import preparar, calcular, guardar
    conteo=Counter();preparadas=[]
    try:
        _opciones(vista['opciones'],session,organizacion_id)
        for fila in sorted(filas,key=lambda f:f['sku']):
            if fila['accion']=='excluido':
                if fila['sku'] not in EXCLUIDOS: raise ValueError('La exclusión no corresponde al lote.')
                conteo['excluidos']+=1;continue
            existentes=session.query(Producto).filter(Producto.organizacion_id==organizacion_id,Producto.sku.ilike(fila['sku'])).all()
            if fila['accion']=='conservar':
                if len(existentes)!=1 or existentes[0].id!=fila['producto_id']: raise ValueError('Cambió un producto conservado. Volvé a previsualizar el lote.')
                conteo['conservados']+=1;continue
            if fila['accion']!='crear': raise ValueError('La acción de la fila no es válida.')
            if existentes: raise ValueError(f'{fila["sku"]} ya existe. Volvé a previsualizar para conservarlo sin duplicar.')
            plan=preparar(fila['produccion'],session,organizacion_id)
            if calcular(plan,organizacion_id)['huella']!=fila['huella_costo']: raise ValueError('Cambiaron los costos desde la vista previa. Volvé a analizar el archivo.')
            preparadas.append((fila,plan,leer_formulario_completo(fila['adicionales'],[])))
        for fila,plan,adicionales in preparadas:
            producto=guardar_ficha(fila['datos'],Producto=Producto,db_session=session,organizacion_id=organizacion_id,huella_esperada='nuevo',commit=False)
            guardar_datos_generales(producto,adicionales,session,usuario=usuario)
            guardar(producto,plan,session,usuario,fila['huella_costo'])
            conteo['creados']+=1
        session.commit()
    except Exception:
        session.rollback();raise
    return {c:conteo[c] for c in ('creados','conservados','excluidos')}
