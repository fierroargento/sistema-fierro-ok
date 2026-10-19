"""Previsualización firmada y aplicación atómica del mismo contrato manual."""
from collections import Counter
from io import BytesIO, StringIO
import csv
import unicodedata

from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer

from services.productos_ficha_maestra import (
    CAMPOS_BASE, CAMPOS_LOGISTICA, construir_filas, guardar_ficha,
    normalizar_ficha, validar_vinculos,
)
from services.producto_habilitaciones import CAMPOS_FICHA
from services.productos_combo_maestro import preparar_composicion

MAX_FILAS = 500
CAMPOS = ("sku", "nombre", "linea", "descripcion", "tipo", "unidad_medida", "activo", "no_stockeable", "permitir_stock_negativo",
    "habilita_compra", "habilita_venta", "habilita_insumo", "habilita_componente", "insumo_id", "componentes_combo", "unidad_combo_id") + CAMPOS_LOGISTICA
ALIASES = {"id_producto": "sku", "codigo": "sku", "codigo_producto": "sku",
    "activo_si_no": "activo", "tipo_de_producto": "tipo", "unidad": "unidad_medida",
    "compra": "habilita_compra", "venta": "habilita_venta", "insumo": "habilita_insumo",
    "componente": "habilita_componente", "peso_gr": "peso_gr", "peso_g": "peso_gr",
    "permite_via_cargo": "permite_via_cargo", "observacion_logistica": "observacion_logistica"}


def _columna(valor):
    limpio = "".join(c for c in unicodedata.normalize("NFD", str(valor or "")) if unicodedata.category(c) != "Mn").lower()
    return "_".join(limpio.replace("(", " ").replace(")", " ").replace("/", " ").replace("_", " ").split())


def leer_lote(archivo):
    contenido = archivo.read(2 * 1024 * 1024 + 1)
    if len(contenido) > 2 * 1024 * 1024:
        raise ValueError("El archivo supera 2 MB. Dividilo en lotes de hasta 500 productos.")
    nombre = str(archivo.filename or "").lower()
    if nombre.endswith(".csv"):
        texto = contenido.decode("utf-8-sig")
        try:
            dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
        except csv.Error:
            dialecto = csv.excel
        fuente = csv.reader(StringIO(texto), dialecto)
        libro = None
    elif nombre.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        libro = load_workbook(BytesIO(contenido), read_only=True, data_only=False, keep_links=False)
        hoja = libro["Productos"] if "Productos" in libro.sheetnames else libro.worksheets[0]
        fuente = hoja.iter_rows(values_only=True)
    else:
        raise ValueError("Elegí un archivo XLSX, XLSM o CSV. Para XLS antiguo, guardalo primero como XLSX.")
    try:
        encabezados = next(fuente, None)
        if not encabezados:
            raise ValueError("El archivo está vacío.")
        campos = [ALIASES.get(_columna(c), _columna(c)) for c in encabezados]
        reconocidos = [c for c in campos if c in CAMPOS]
        if "sku" not in reconocidos:
            raise ValueError("Falta la columna SKU o ID Producto.")
        if len(reconocidos) != len(set(reconocidos)):
            raise ValueError("Hay columnas repetidas para el mismo campo. Conservá una sola.")
        filas = []
        for numero, valores in enumerate(fuente, 2):
            if numero > MAX_FILAS + 2:
                raise ValueError("El archivo supera 500 filas. Dividilo en lotes más pequeños.")
            datos = {c: valores[i] for i, c in enumerate(campos) if c in CAMPOS and i < len(valores) and valores[i] not in (None, "")}
            if not datos:
                continue
            if len(filas) >= MAX_FILAS:
                raise ValueError("El archivo supera 500 productos.")
            filas.append({"numero": numero, "datos": datos})
        if not filas:
            raise ValueError("No se encontraron productos.")
        return filas
    finally:
        if libro is not None:
            libro.close()


def previsualizar_lote(filas, *, Producto, db_session, organizacion_id):
    productos = db_session.query(Producto).filter_by(organizacion_id=organizacion_id).all()
    listo, actuales = construir_filas(productos, db_session, organizacion_id)
    if not listo:
        raise ValueError("La ficha ampliada todavía no está preparada.")
    por_sku = {}
    for fila in actuales:
        sku = fila["datos"]["sku"].strip().upper()
        if sku in por_sku:
            raise ValueError(f"El maestro contiene un SKU repetido: {sku}. Revisalo antes de importar.")
        por_sku[sku] = fila
    modelos = {p.id: p for p in productos}
    repetidos = Counter(str(f["datos"].get("sku") or "").strip().upper() for f in filas)
    resultado = []
    for original in filas:
        entrada = original["datos"]
        sku = str(entrada.get("sku") or "").strip().upper()
        anterior = por_sku.get(sku)
        base = dict(anterior["datos"]) if anterior else {
            "linea": "", "descripcion": "", "insumo_id": None,
            "permite_correo": False, "permite_via_cargo": False,
            "requiere_revision_logistica": True, "observacion_logistica": "",
        }
        # En un producto aún sin configurar se exige elegir los usos, no
        # convertir las casillas vacías de la vista en un default del lote.
        if anterior and not anterior["configurado"]:
            for c in ("activo", "habilita_compra", "habilita_venta", "habilita_insumo", "habilita_componente"):
                base.pop(c, None)
        base.update(entrada)
        errores, cambios, datos = [], [], None
        try:
            if repetidos[sku] > 1:
                raise ValueError(f"SKU duplicado {sku}: revisá todas sus filas; no se eligió ninguna.")
            if any((isinstance(v, str) and v.startswith("="))
                    or not isinstance(v, (str, int, float, bool, type(None)))
                    for v in entrada.values()):
                raise ValueError("La fila contiene fórmulas. Pegá los datos verificados como valores.")
            datos = normalizar_ficha(base)
            validar_vinculos(modelos.get(anterior["id"]) if anterior else None, datos, db_session, organizacion_id)
            preparar_composicion(modelos.get(anterior["id"]) if anterior else None, datos, db_session, organizacion_id)
            for campo in CAMPOS:
                antes = anterior["datos"].get(campo) if anterior else None
                if campo == "sku" and antes is not None:
                    antes = antes.strip().upper()
                if not anterior or not anterior["configurado"] or antes != datos.get(campo):
                    if campo in entrada or not anterior:
                        cambios.append({"campo": campo, "antes": antes, "despues": datos.get(campo)})
            accion = "crear" if anterior is None else "actualizar" if cambios or not anterior["configurado"] else "sin_cambios"
        except (ValueError, TypeError) as error:
            errores.append(str(error))
            accion = "rechazado"
        resultado.append({"numero": original["numero"], "sku": sku, "datos": datos,
            "producto_id": anterior["id"] if anterior else None,
            "huella": anterior["huella"] if anterior else "nuevo",
            "accion": accion, "cambios": cambios, "errores": errores})
    vinculados = [r["datos"]["insumo_id"] for r in resultado if r["datos"] and r["datos"]["insumo_id"] is not None]
    duplicados = {i for i, n in Counter(vinculados).items() if n > 1}
    for fila in resultado:
        if fila["datos"] and fila["datos"]["insumo_id"] in duplicados:
            fila["accion"] = "rechazado"
            fila["errores"].append("El lote vincula el mismo insumo a varios productos.")
    return resultado


def firmar_lote(vista, organizacion_id, usuario_id, secret_key):
    return URLSafeTimedSerializer(secret_key, salt="maestro-productos-v1").dumps(
        {"organizacion_id": organizacion_id, "usuario_id": usuario_id, "filas": vista})


def aplicar_lote(token, *, organizacion_id, usuario_id, secret_key, Producto, db_session):
    if len(token) > 500000:
        raise ValueError("La previsualización no es válida.")
    try:
        lote = URLSafeTimedSerializer(secret_key, salt="maestro-productos-v1").loads(token, max_age=900)
    except (BadSignature, SignatureExpired) as error:
        raise ValueError("La previsualización venció o cambió. Volvé a subir el archivo.") from error
    if lote["organizacion_id"] != organizacion_id or lote["usuario_id"] != usuario_id:
        raise ValueError("El lote no pertenece al usuario y organización activos.")
    if any(f["errores"] for f in lote["filas"]):
        raise ValueError("Corregí las filas rechazadas antes de confirmar. No se aplicó el lote.")
    conteo = {"crear": 0, "actualizar": 0, "sin_cambios": 0}
    try:
        # Orden estable de bloqueo y una única transacción para todo el lote.
        for fila in sorted(lote["filas"], key=lambda f: f["producto_id"] or 0):
            if fila["accion"] == "sin_cambios":
                # Incluso estas filas se releen para rechazar un lote obsoleto.
                from models.producto_ficha_maestra import ProductoFichaMaestra
                from services.productos_ficha_maestra import huella
                producto = db_session.query(Producto).filter_by(id=fila["producto_id"], organizacion_id=organizacion_id).with_for_update().populate_existing().first()
                ficha = db_session.query(ProductoFichaMaestra).filter_by(producto_id=fila["producto_id"], organizacion_id=organizacion_id).populate_existing().first()
                if producto is None or huella(producto, ficha) != fila["huella"]:
                    raise ValueError("El lote quedó desactualizado. Volvé a previsualizarlo.")
            else:
                guardar_ficha(fila["datos"], Producto=Producto, db_session=db_session,
                    organizacion_id=organizacion_id, producto_id=fila["producto_id"],
                    huella_esperada=fila["huella"], commit=False)
            conteo[fila["accion"]] += 1
        db_session.commit()
    except Exception:
        db_session.rollback()
        raise
    return conteo


def plantilla_csv():
    salida = StringIO(newline="")
    escritor = csv.writer(salida, delimiter=";")
    escritor.writerow(CAMPOS)
    # Ejemplo completo y desconectado, sin permisos logísticos por defecto.
    ejemplo = {"sku": "UAT-PRODUCTO-001", "nombre": "Producto de prueba", "linea": "Pruebas",
        "descripcion": "Producto de prueba", "tipo": "simple", "unidad_medida": "unidad",
        "activo": "si", "no_stockeable": "no", "permitir_stock_negativo": "no", "habilita_compra": "si", "habilita_venta": "si",
        "habilita_insumo": "si", "habilita_componente": "si", "permite_correo": "no",
        "permite_via_cargo": "no", "requiere_revision_logistica": "si"}
    escritor.writerow([ejemplo.get(c, "") for c in CAMPOS])
    return "\ufeff" + salida.getvalue()
