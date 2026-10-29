"""Adaptador explícito Codimat: números originales, sin ejecutar fórmulas del XLSX."""
from decimal import Decimal,InvalidOperation,ROUND_HALF_UP
from io import BytesIO
import re
from zipfile import ZipFile,BadZipFile
from models.listas_proveedores_compra import ConfiguracionListaProveedor as Config

DESCUENTOS={'1955':40,'1957':40,'31020':46,'31030':46,'1867':40,'3951':40,'1984':40,'1974':40,'21572':38,'21573':38,'21574':38,'21575':38,'2393':46}
ENCABEZADOS=['Tipo Unidad','Cant Unidad','Código','Descripción','IVA','34%+7%','34%','BonifC','$ C/IVA(PESOS)','$ S/IVA(PESOS)']

def configuracion(org,unidad,proveedor,session,bloquear=False):
    q=session.query(Config).filter_by(organizacion_id=org,unidad_negocio_id=unidad,proveedor_id=proveedor)
    return (q.with_for_update() if bloquear else q).populate_existing().first()

def reglas_actuales(org,unidad,proveedor,session):
    c=configuracion(org,unidad,proveedor,session)
    return c.reglas if c else [dict(codigo=k,descuento=str(v),cantidad=str(12 if k=='1867' else 2 if k=='2393' else 6)) for k,v in DESCUENTOS.items()]

def decimal(v,etiqueta):
    try:
        if isinstance(v,bool) or str(v).strip().startswith('='):raise ValueError()
        texto=str(v).strip()
        if not re.fullmatch(r'[+-]?(?:[0-9]+(?:[.,][0-9]*)?|[.,][0-9]+)%?',texto):raise ValueError()
        n=Decimal(texto.removesuffix('%').replace(',','.'))
        if not n.is_finite():raise ValueError()
        return n
    except (InvalidOperation,ValueError):raise ValueError(etiqueta+' inválido.') from None

def reglas(codigos,descuentos,cantidades):
    if not(len(codigos)==len(descuentos)==len(cantidades)) or len(codigos)>500:raise ValueError('Revisá las bonificaciones por código.')
    salida=[];vistos=set()
    for cod,d,c in zip(codigos,descuentos,cantidades):
        cod=str(cod).strip().upper()
        if not cod and not str(d).strip() and not str(c).strip():continue
        if not cod or len(cod)>120 or cod in vistos:raise ValueError('Completá códigos únicos para las bonificaciones.')
        n=decimal(d,'Bonificación')
        if not 0<=n<=100:raise ValueError('La bonificación debe estar entre 0 y 100%.')
        cantidad=''
        if str(c).strip():
            q=decimal(c,'Cantidad por bulto')
            if not 0<q<=Decimal('999999999999'):raise ValueError('La cantidad por bulto debe ser positiva.')
            cantidad=str(q)
        vistos.add(cod);salida.append(dict(codigo=cod,descuento=str(n),cantidad=cantidad))
    if not salida:raise ValueError('Completá al menos una bonificación por código.')
    return salida

def revisar(archivo,reglas,p,org,unidad,session):
    from services import listas_proveedores_compra as listas
    from models.insumo_presentacion_compra import InsumoPresentacionCompra as IP
    from models.listas_proveedores_compra import CotizacionProveedorCompra as C
    listas.listo(session);listas.proveedor_valido(p,org,unidad,session)
    if not str(archivo.filename or '').lower().endswith('.xlsx'):raise ValueError('Codimat: subí el XLSX original.')
    contenido=archivo.read(10*1024*1024+1)
    if len(contenido)>10*1024*1024:raise ValueError('La planilla supera 10 MB.')
    from openpyxl import load_workbook
    try:
        with ZipFile(BytesIO(contenido)) as z:
            if len(z.infolist())>2000 or sum(i.file_size for i in z.infolist())>100*1024*1024:raise ValueError('La planilla descomprimida supera el límite.')
        w=load_workbook(BytesIO(contenido),read_only=True,data_only=False,keep_links=False)
    except (BadZipFile,KeyError,OSError):raise ValueError('No se pudo leer el XLSX de Codimat.') from None
    cat,materiales=listas.catalogo(org,unidad,p.id,session)
    por_codigo={r['codigo']:r for r in reglas};filas=[];ignoradas=0;encontrados=set()
    ids=[r['vinculo_id'] for r in cat.values() if r['vinculo_id']]
    previas={c.vinculo_id:c for c in session.query(C).filter(C.vinculo_id.in_(ids),C.moneda=='ARS',C.vigente.is_(True))}
    presentaciones={x.insumo_id:x for x in session.query(IP).filter(IP.organizacion_id==org,IP.insumo_id.in_(list(materiales)))}
    try:
        if 'Lista de Precios' not in w:raise ValueError('Codimat: falta la hoja Lista de Precios.')
        hoja=w['Lista de Precios']
        # El XLSX original declara A1 aunque contiene miles de filas y diez columnas.
        # Ignorar esa dimensión incorrecta y acotar la lectura real, sin confiar en ella.
        hoja.reset_dimensions()
        cabecera=None
        for row in hoja.iter_rows(min_row=1,max_row=10,max_col=10):
            if [c.value for c in row[:10]]==ENCABEZADOS:cabecera=row[0].row;break
        if cabecera is None:raise ValueError('Cambió el formato de Codimat. Revisá los encabezados antes de importar.')
        for numero,row in enumerate(hoja.iter_rows(min_row=cabecera+1,max_col=10),cabecera+1):
            if numero>50006:raise ValueError('La lista supera 50.000 filas.')
            v=[c.value for c in row];cod=str(v[2] or '').strip().upper()
            if not cod:continue
            ref=cat.get(cod)
            if not ref:ignoradas+=1;continue
            if len(filas)>=500:raise ValueError('La lista supera 500 coincidencias.')
            encontrados.add(cod);errores=[];datos={};calculo={}
            try:
                regla=por_codigo.get(cod)
                if regla is None:raise ValueError('Falta configurar la bonificación de este código. No se asume 0%.')
                bruto=decimal(v[6],'Precio columna 34%');iva=decimal(v[4],'Factor IVA');bonif=decimal(regla['descuento'],'Bonificación')
                if not 0<=bruto<=Decimal('999999999999.99') or not 1<=iva<=2:raise ValueError('Precio o factor IVA fuera de rango.')
                if any(row[j].data_type=='f' for j in (2,4,6,7)):raise ValueError('Código, IVA, precio 34% y BonifC deben contener valores, no fórmulas.')
                # BonifC se reemplaza por la acordada, no se suma ni se aplica dos veces.
                bonif_archivo=decimal(v[7] if v[7] is not None else 0,'BonifC del archivo')
                if not 0<=bonif_archivo<=100:raise ValueError('BonifC del archivo fuera de rango.')
                neto=bruto*(1-bonif/100)/iva
                q=regla['cantidad'];previa=previas.get(ref['vinculo_id']);presentacion=presentaciones.get(ref['insumo_id'])
                if not q:q=str(previa.cantidad_bulto) if previa and previa.cantidad_bulto else str(presentacion.cantidad_bulto) if presentacion and presentacion.cantidad_bulto else ''
                if not q:raise ValueError('Completá cantidad de consumo por bulto. Cant Unidad del Excel puede ser peso y no se utiliza para esta conversión.')
                calculo=dict(formato='Codimat',precio_34_con_iva=str(bruto),bonificacion=str(bonif),bonificacion_archivo=str(bonif_archivo),factor_iva=str(iva),bulto_con_iva=str((bruto*(1-bonif/100)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)),bulto_sin_iva=str(neto.quantize(Decimal('.01'),rounding=ROUND_HALF_UP)),cantidad=str(q),descripcion=str(v[3] or '')[:300],tipo_unidad=str(v[0] or ''),cant_unidad_archivo=str(v[1] or ''))
                datos=dict(precio_proveedor=calculo['bulto_sin_iva'],base_precio='bulto',cantidad_bulto=q,moneda_proveedor='ARS',medidas_origen=previa.medidas_origen if previa else (presentacion.medidas_origen or '') if presentacion else '')
            except ValueError as e:errores=[str(e)]
            datos.update(codigo_proveedor=cod,_codimat=calculo,_error_codimat=errores)
            filas.append(dict(numero=numero,valores=datos,formulas=[]))
    finally:w.close()
    vista=listas.revisar(filas,list(listas.CAMPOS),p,org,unidad,session)
    vista['ignoradas']=ignoradas
    c=configuracion(org,unidad,p.id,session)
    vista['codimat']=dict(reglas=reglas,config_anterior=c.reglas if c else None,codigos_ausentes=sorted(set(por_codigo)-encontrados))
    return vista
