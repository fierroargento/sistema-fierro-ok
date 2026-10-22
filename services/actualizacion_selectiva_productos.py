"""Plantilla vacía y parches explícitos del maestro, sin reemplazar campos ajenos."""
from collections import Counter, defaultdict
from decimal import Decimal
from hashlib import sha256
from io import BytesIO, StringIO
from zipfile import ZipFile, BadZipFile
import csv
import json
from itsdangerous import URLSafeTimedSerializer, BadSignature
from sqlalchemy import inspect

from services import planilla_productos as completo
from services.productos_ficha_maestra import CAMPOS_LOGISTICA, normalizar_logistica, texto

SALT = 'productos-campos-seleccionados-v1'
MAX_FILAS_LISTA = 50000
MAX_COINCIDENCIAS = 500


class LimiteCoincidencias(ValueError):
    pass
RELACIONES = ('codigos_barras', 'sucursales', 'depositos')
PRECIOS = ('precio_proveedor', 'base_precio', 'cantidad_bulto', 'moneda_proveedor', 'medidas_origen')
CAMPOS = tuple(c for c in completo.BASE + completo.GENERAL if c not in ('sku', 'insumo_id')) + RELACIONES + PRECIOS
ETIQUETAS = dict(completo.ETIQUETAS, precio_proveedor='Precio proveedor sin IVA',
    base_precio='Cotización proveedor', cantidad_bulto='Cantidad por bulto',
    moneda_proveedor='Moneda proveedor', medidas_origen='Medidas en origen',
    codigo_proveedor='Código proveedor', codigos_barras='Códigos de barras',
    sucursales='Sucursales habilitadas', depositos='Depósitos habilitados')
CLAVES = ('sku', 'codigo_proveedor')
GRUPOS = {
 'Identificación y clasificación': ('nombre','linea','descripcion','tipo','unidad_medida','rubro','subrubro','marca','codigos_barras'),
 'Usos y stock': ('activo','no_stockeable','permitir_stock_negativo','habilita_compra','habilita_venta','habilita_insumo','habilita_componente','sucursales','depositos','sucursales_definidas','depositos_definidos'),
 'Peso y envíos': CAMPOS_LOGISTICA,
 'Costo simple e impuestos': ('precio_costo','moneda_costo','iva_pct','iibb_pct','utilidad_pct','utilidad_base'),
 'Lista de proveedor': PRECIOS,
}


def etiqueta(c):return ETIQUETAS.get(c, completo.etiqueta(c))


def plantilla():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    w=Workbook();s=w.active;s.title='Productos'
    columnas=CLAVES+CAMPOS;s.append([etiqueta(c) for c in columnas])
    for i,c in enumerate(columnas,1):
        letra=s.cell(1,i).column_letter
        s.column_dimensions[letra].width=38 if c in ('nombre','descripcion','observacion_logistica') else 25
        if c in CLAVES+('codigos_barras',):s.column_dimensions[letra].number_format='@'
        s.cell(1,i).font=Font(bold=True,color='FFFFFF')
        s.cell(1,i).fill=PatternFill('solid',fgColor='17566D')
        s.cell(1,i).alignment=Alignment(wrap_text=True,vertical='center')
        # Identificadores de texto: conservar ceros iniciales al pegar códigos.
        for n in range(2,502):
            s.cell(n,i).number_format='@' if c in CLAVES+('codigos_barras',) else 'General'
        opciones=('Sí,No' if c in completo.BOOLEANOS else 'simple,produccion,combo' if c=='tipo' else
                  'ARS,USD' if c in ('moneda_costo','moneda_proveedor') else 'unitario,bulto' if c=='base_precio' else
                  'costo,precio_venta' if c=='utilidad_base' else None)
        if opciones:
            dv=DataValidation(type='list',formula1='"'+opciones+'"');dv.showErrorMessage=True;dv.error='Elegí una opción de la lista.';s.add_data_validation(dv);dv.add(f'{letra}2:{letra}{MAX_FILAS_LISTA+1}')
    s.row_dimensions[1].height=36;s.freeze_panes='C2';s.auto_filter.ref=f'A1:{s.cell(1,len(columnas)).column_letter}{MAX_FILAS_LISTA+1}';s.sheet_view.showGridLines=False
    a=w.create_sheet('Cómo usar')
    for fila in [
        ['Actualizar sólo los campos elegidos'],
        ['1. Pegá SKU y los datos que querés modificar. Dejá vacíos los demás. No cambies los encabezados.'],
        ['2. Alternativa: elegí un proveedor en Fierro y pegá Código proveedor. Si incluís SKU, ambos deben coincidir.'],
        ['3. Al subir, tildá sólo los campos que querés actualizar. Los no seleccionados y los vacíos se conservan; cero y No son valores.'],
        ['4. Revisá antes/después y confirmá. Esta carga actualiza productos existentes; no crea SKU nuevos.'],
        ['Peso en gramos; medidas en cm. Para kg convertí a gramos antes de pegar.'],
        ['Sucursales y depósitos: códigos de las ubicaciones de Fierro separados con |. Barras: códigos separados con |. Las listas seleccionadas y completas reemplazan esa lista; vacías conservan.'],
        ['Precios sin IVA: Precio proveedor sin IVA se interpreta según Cotización proveedor (unitario o bulto). Bulto exige cantidad. Moneda ARS o USD.'],
        ['El precio proveedor actualiza la fuente de insumo vinculada, con historial; para simples también su costo maestro. Para producción o combo sin fuente vinculada se rechaza el precio manual.'],
        ['Costo simple sin IVA es exclusivo de productos simples. Producción y combo construyen su costo desde componentes y tareas.'],
        ['Tipo se conserva en fichas configuradas. Una ficha antigua puede completarse como simple; producción o combo requieren su composición desde el engranaje.'],
        ['Imágenes, vínculos de insumo, recetas de fabricación/combo, stock e historial no se reemplazan desde esta plantilla. Para revisar piezas/tareas usá Exportar para editar.'],
        ['Si no hay SKU o Código proveedor no se considera la fila. Hasta 50.000 filas de lista, 500 coincidencias y 10 MB, XLSX o CSV. Las filas sin coincidencia se ignoran automáticamente. Fórmulas en campos seleccionados: pegar valores.'],
    ]:a.append(fila)
    a.column_dimensions['A'].width=115
    for row in a:
        row[0].alignment=Alignment(wrap_text=True,vertical='top')
        a.row_dimensions[row[0].row].height=45
    out=BytesIO();w.save(out);out.seek(0);return out


def leer(archivo, seleccion):
    campos=list(dict.fromkeys(seleccion))
    if not campos or any(c not in CAMPOS for c in campos):
        raise ValueError('Tildá al menos un campo válido para actualizar.')
    contenido=archivo.read(10*1024*1024+1)
    if len(contenido)>10*1024*1024:raise ValueError('La planilla supera 10 MB.')
    nombre=str(archivo.filename or '').lower();w=None
    if nombre.endswith('.xlsx'):
        try:
            with ZipFile(BytesIO(contenido)) as z:
                if len(z.infolist())>2000 or sum(i.file_size for i in z.infolist())>100*1024*1024:raise ValueError('La planilla descomprimida supera el límite.')
            from openpyxl import load_workbook
            w=load_workbook(BytesIO(contenido),read_only=True,data_only=False,keep_links=False)
            if 'Control' in w:raise ValueError('Esta es una exportación completa. Subila en Actualizar desde planilla Fierro o usá la plantilla vacía.')
            if 'Productos' not in w:raise ValueError('Falta la hoja Productos.')
            if w['Productos'].max_row>MAX_FILAS_LISTA+1 or w['Productos'].max_column>100:raise ValueError('La hoja supera los límites permitidos.')
            fuente=w['Productos'].iter_rows();enc=[c.value for c in next(fuente,())]
        except (BadZipFile,KeyError):raise ValueError('No se pudo leer el XLSX.') from None
    elif nombre.endswith('.csv'):
        try:txt=contenido.decode('utf-8-sig')
        except UnicodeDecodeError:txt=contenido.decode('cp1252')
        try:dialecto=csv.Sniffer().sniff(txt[:4096],delimiters=';,\t')
        except csv.Error:dialecto=None
        fuente=csv.reader(StringIO(txt),dialecto) if dialecto else csv.reader(StringIO(txt),delimiter=';')
        enc=next(fuente,[])
    else:raise ValueError('Usá la plantilla vacía guardada como XLSX o CSV.')
    try:
        nombres={etiqueta(c):c for c in CLAVES+CAMPOS}
        columnas=[nombres.get(str(c).strip(),str(c).strip() if str(c).strip() in CLAVES+CAMPOS else None) for c in enc]
        elegidos=set(campos)|set(CLAVES)
        relevantes=[c for c in columnas if c in elegidos]
        if len(relevantes)!=len(set(relevantes)):raise ValueError('Hay encabezados repetidos en los campos seleccionados.')
        if not any(c in columnas for c in CLAVES):raise ValueError('Falta SKU o Código proveedor.')
        faltantes=[etiqueta(c) for c in campos if c not in columnas]
        if faltantes:raise ValueError('Faltan las columnas seleccionadas: '+', '.join(faltantes)+'.')
        filas=[]
        for numero,row in enumerate(fuente,2):
            if numero>MAX_FILAS_LISTA+1:raise ValueError('La lista supera 50.000 filas.')
            valores={};formulas=[]
            for i,c in enumerate(columnas):
                if c not in elegidos or i>=len(row):continue
                celda=row[i];v=celda.value if w else celda
                if c in CLAVES and w and celda.data_type=='f':
                    formulas.append(c)
                if v is None or str(v).strip()=='':continue
                if c in campos and ((w and celda.data_type=='f') or isinstance(v,str) and v.startswith('=')):
                    formulas.append(c)
                if isinstance(v,(int,float)) and not isinstance(v,bool):v=format(Decimal(str(v)),'f')
                valores[c]=str(v).strip() if not isinstance(v,bool) else v
            if not any(valores.get(c) for c in CLAVES):continue
            if len(filas)>=MAX_FILAS_LISTA:raise ValueError('La lista supera 50.000 artículos.')
            filas.append({'numero':numero,'valores':valores,'formulas':formulas})
        if not filas:raise ValueError('Pegá al menos un SKU o Código proveedor y sus datos.')
        return filas,campos
    finally:
        if w:w.close()


def _identidades(session,org):
    from models.producto_datos_generales import ProductoProveedor
    from models.compras import ProveedorCompra
    proveedores={p.id:p for p in session.query(ProveedorCompra).filter_by(organizacion_id=org)}
    relaciones=list(session.query(ProductoProveedor).filter_by(organizacion_id=org))
    return proveedores,relaciones


def _lista(valor):return [v.strip() for v in str(valor).split('|') if v.strip()]


def _datos_precio(entrada,actual,proveedor_id,session,org,cache):
    from models.fuentes_costo_productivo import InsumoProductivo, InsumoPrecioVersion
    from services import insumos_presentacion as bultos
    insumo_id=actual['datos']['insumo_id']
    if 'insumos' not in cache:
        cache['insumos']={i.id:i for i in session.query(InsumoProductivo).filter_by(organizacion_id=org)}
        cache['precios_disponibles']=inspect(session.connection()).has_table(InsumoPrecioVersion.__tablename__)
        cache['vigentes']={(v.insumo_id,v.moneda):v for v in session.query(InsumoPrecioVersion).join(InsumoProductivo,InsumoProductivo.id==InsumoPrecioVersion.insumo_id).filter(InsumoProductivo.organizacion_id==org,InsumoPrecioVersion.vigente.is_(True))} if cache['precios_disponibles'] else {}
        cache['bultos_disponibles']=bultos.disponible(session)
        cache['presentaciones']={}
        from models.insumo_presentacion_compra import InsumoPrecioBultoVersion
        cache['historicos']={v.precio_version_id:v for v in session.query(InsumoPrecioBultoVersion).filter_by(organizacion_id=org)} if cache['bultos_disponibles'] else {}
    insumo=cache['insumos'].get(insumo_id)
    if insumo_id and not cache['precios_disponibles']:raise ValueError('Falta preparar los precios de insumos en staging.')
    if insumo_id and insumo is None:raise ValueError('La fuente de insumo no pertenece a esta organización.')
    if not insumo and actual['datos']['tipo']!='simple':
        raise ValueError('Precio proveedor: vinculá la fuente de insumo desde el engranaje. El costo de fabricación/combo no se carga a mano.')
    monedas=[m for (i,m) in cache['vigentes'] if i==insumo_id]
    if not entrada.get('moneda_proveedor') and len(monedas)>1:
        raise ValueError('La fuente tiene precios en varias monedas. Completá y seleccioná Moneda proveedor.')
    moneda=str(entrada.get('moneda_proveedor') or (monedas[0] if monedas else actual['adicionales'].get('moneda_costo') or 'ARS')).upper()
    if moneda not in ('ARS','USD'):raise ValueError('Moneda proveedor: elegí ARS o USD.')
    if insumo and insumo.unidad_negocio_id not in cache['presentaciones']:
        cache['presentaciones'][insumo.unidad_negocio_id]=bultos.leer(session,org,insumo.unidad_negocio_id)
    presentaciones=cache['presentaciones'].get(insumo.unidad_negocio_id,{}) if insumo else {}
    previa=presentaciones.get(insumo_id,{})
    cantidad=entrada.get('cantidad_bulto') or previa.get('cantidad_bulto') or ''
    # La base es explícita: no interpretar un importe por bulto como unitario.
    base=str(entrada.get('base_precio') or '').lower()
    if base not in ('unitario','bulto'):raise ValueError('Completá y seleccioná Cotización proveedor: unitario o bulto.')
    if not insumo and (base=='bulto' or entrada.get('cantidad_bulto') or entrada.get('medidas_origen')):
        raise ValueError('Para guardar bulto y medidas de origen, vinculá la fuente de insumo desde el engranaje. Sin fuente vinculada podés actualizar el costo simple por unidad.')
    datos={'base_precio':base,'cantidad_bulto':str(cantidad),'medidas_origen':entrada.get('medidas_origen') or previa.get('medidas_origen',''),
           'precio_bulto' if base=='bulto' else 'precio_unitario':entrada['precio_proveedor']}
    plan=bultos.preparar_precio(datos)
    if insumo and (plan['cantidad_bulto'] or plan['bulto_centavos'] is not None) and not cache['bultos_disponibles']:
        raise ValueError('Los bultos del insumo todavía no están preparados.')
    vigente=cache['vigentes'].get((insumo_id,moneda)) if insumo else None
    firma={'insumo_id':insumo_id,'unidad':insumo.unidad_medida if insumo else None,'version_id':vigente.id if vigente else None,
           'precio':vigente.precio_unitario_centavos if vigente else None,'presentacion':completo._normal({k:v for k,v in previa.items() if k!='historial'})}
    # El historial de bulto completo puede tener claves int; se firma con JSON normalizado.
    hist=cache['historicos'].get(vigente.id) if vigente else None
    cantidad_plan=Decimal(plan['cantidad_bulto']) if plan['cantidad_bulto'] else None
    cambia=bool(insumo and (vigente is None or vigente.precio_unitario_centavos!=plan['unitario_centavos']
        or (cantidad_plan is not None and (hist is None or hist.cantidad_bulto!=cantidad_plan or hist.precio_bulto_centavos!=plan['bulto_centavos'] or hist.base_precio!=base))
        or (entrada.get('medidas_origen') and entrada['medidas_origen']!=previa.get('medidas_origen'))))
    return {'cambia':cambia,'proveedor_id':proveedor_id,'insumo_id':insumo_id,'moneda':moneda,'datos':datos,'plan':plan,
            'huella':sha256(json.dumps(firma,sort_keys=True,default=str).encode()).hexdigest(),
            'precio_anterior':vigente.precio_unitario_centavos if vigente else None}


def previsualizar(filas,campos,*,proveedor_id,Producto,session,organizacion_id):
    if not campos or any(c not in CAMPOS for c in campos):raise ValueError('Elegí campos válidos.')
    actuales=completo.actuales(Producto,session,organizacion_id)
    proveedores,relaciones=_identidades(session,organizacion_id)
    if proveedor_id and not str(proveedor_id).isdigit():raise ValueError('Elegí un proveedor válido.')
    proveedor_id=int(proveedor_id) if str(proveedor_id or '').isdigit() else None
    if proveedor_id and proveedor_id not in proveedores:raise ValueError('El proveedor no pertenece a esta organización.')
    if not proveedor_id and any(r['valores'].get('codigo_proveedor') for r in filas):
        raise ValueError('Elegí el proveedor de esta lista para buscar por Código proveedor.')
    por_id={f['id']:f for f in actuales.values()}
    por_codigo=defaultdict(list);vinculados=set()
    for x in relaciones:
        por_codigo[(x.proveedor_id,str(x.codigo_clave or '').strip().upper())].append(x)
        vinculados.add((x.proveedor_id,x.producto_id))
    vista=[];tablas={h:[] for h in completo.COLUMNAS};legados=[];precios=[];cache={}
    entrada_coincidente=[];ignoradas=0;muestras_ignoradas=[]
    def ignorar(r,sku,codigo,motivo):
        nonlocal ignoradas
        ignoradas+=1
        if len(muestras_ignoradas)<20:muestras_ignoradas.append({'numero':r['numero'],'sku':sku,'codigo_proveedor':codigo,'motivo':motivo})
    for r in filas:
        entrada={c:v for c,v in r['valores'].items() if c in set(campos)|set(CLAVES)}
        sku=str(entrada.get('sku') or '').upper()
        f={'numero':r['numero'],'sku':sku,'accion':'rechazado','errores':[],'cambios':[],'huella':'','producto_id':None}
        try:
            codigo=str(entrada.get('codigo_proveedor') or '').strip().upper()
            actual=actuales.get(sku) if sku else None
            if codigo:
                if proveedor_id is None:raise ValueError('Elegí el proveedor de esta lista para buscar por Código proveedor.')
                matches=por_codigo.get((proveedor_id,codigo),[])
                if not matches:
                    ignorar(r,sku,codigo,'Código no cargado para el proveedor elegido.');continue
                if len(matches)!=1:raise ValueError('Código proveedor ambiguo: identifica varios productos existentes. Revisá sus vínculos.')
                identificado=por_id.get(matches[0].producto_id)
                if not identificado or (sku and actual is None):
                    ignorar(r,sku,codigo,'Sin coincidencia de SKU y código para este proveedor.');continue
                if sku and actual['id']!=identificado['id']:raise ValueError('SKU y Código proveedor apuntan a productos distintos. Corregí la fila.')
                actual=identificado;sku=actual['datos']['sku'];f['sku']=sku
            if actual is None:
                ignorar(r,sku,codigo,'SKU no cargado en Fierro.');continue
            if proveedor_id and (proveedor_id,actual['id']) not in vinculados:
                ignorar(r,sku,codigo,'Sin vínculo con el proveedor elegido.');continue
            if len(entrada_coincidente)>=MAX_COINCIDENCIAS:
                raise LimiteCoincidencias('La lista tiene más de 500 coincidencias. Actualizá los productos existentes en bloques de hasta 500.')
            if r.get('formulas'):
                raise ValueError('Pegá valores, no fórmulas, en: '+', '.join(etiqueta(c) for c in r['formulas'])+'.')
            f.update(huella=actual['huella'],producto_id=actual['id'])
            entradas={c:v for c,v in entrada.items() if c in campos}
            if any(c in entradas for c in PRECIOS) and 'precio_proveedor' not in entradas:
                raise ValueError('Los datos de cotización se actualizan junto con Precio proveedor sin IVA. Completá ese importe o destildá esos campos.')
            t=completo.tablas_de_filas({sku:actual});producto=t['Productos'][0]
            for c,v in entradas.items():
                if c in completo.BASE+completo.GENERAL:producto[c]=v
            from models.sucursal_operativa import SucursalOperativa
            for c,hoja,key in [('codigos_barras','Barras','codigo'),('sucursales','Ubicaciones','ubicacion_id'),('depositos','Ubicaciones','ubicacion_id')]:
                if c not in entradas:continue
                valores=_lista(entradas[c])
                if not valores:raise ValueError(etiqueta(c)+': completá la lista; vacía conserva lo guardado.')
                if c=='codigos_barras':t[hoja]=[{'sku':sku,'codigo':v} for v in valores]
                else:
                    tipo='sucursal' if c=='sucursales' else 'deposito'
                    from services.productos_ficha_maestra import booleano
                    bandera='sucursales_definidas' if c=='sucursales' else 'depositos_definidos'
                    if not booleano(producto[bandera],bandera):
                        raise ValueError(etiqueta(c)+': seleccioná también '+etiqueta(bandera)+' y completá Sí para aplicar la lista de ubicaciones.')
                    ubicaciones={x.codigo.upper():x.id for x in session.query(SucursalOperativa).filter_by(organizacion_id=organizacion_id)}
                    if any(v.upper() not in ubicaciones for v in valores):raise ValueError(etiqueta(c)+': usá códigos de ubicaciones existentes separados con |.')
                    t[hoja]=[x for x in t[hoja] if x['tipo']!=tipo]+[{'sku':sku,'tipo':tipo,'ubicacion_id':ubicaciones[v.upper()]} for v in valores]
            precio=None
            if 'precio_proveedor' in entradas:
                if proveedor_id is None:raise ValueError('Elegí un proveedor para actualizar su lista de precios.')
                precio=_datos_precio(entradas,actual,proveedor_id,session,organizacion_id,cache)
                if actual['datos']['tipo']=='simple':
                    from services.productos_costo_maestro import _decimal
                    if 'precio_costo' in entradas and _decimal(entradas['precio_costo'],'Costo simple',Decimal('999999999999.99'),2)!=Decimal(precio['plan']['unitario_centavos'])/100:
                        raise ValueError('Costo simple y precio proveedor calculado no coinciden. Elegí un solo origen del precio o corregí ambos valores.')
                    if 'moneda_costo' in entradas and entradas['moneda_costo']!=precio['moneda']:
                        raise ValueError('Moneda del costo y Moneda proveedor no coinciden.')
                    producto['precio_costo']=str(Decimal(precio['plan']['unitario_centavos'])/100)
                    producto['moneda_costo']=precio['moneda']
                if precio['cambia']:f['cambios'].append({'campo':'Precio proveedor unitario sin IVA ('+precio['moneda']+')','antes':str(Decimal(precio['precio_anterior'])/100) if precio['precio_anterior'] is not None else 'Sin precio de insumo',
                                     'despues':str(Decimal(precio['plan']['unitario_centavos'])/100)})
            # Las fichas logísticas anteriores no requieren inventar Tipo/Unidad para cambiar sólo columnas de Producto.
            solo_producto=set(entradas)<=set(CAMPOS_LOGISTICA+('descripcion',)) and not actual['configurado']
            if solo_producto:
                datos={c:producto[c] for c in CAMPOS_LOGISTICA};validados=normalizar_logistica(datos,campos=entradas)
                if 'descripcion' in entradas:validados['descripcion']=texto(entradas['descripcion'],'descripción',255)
                cambios={c:validados[c] for c in entradas if validados[c]!=actual['datos'].get(c)}
                f['cambios']=[{'campo':etiqueta(c),'antes':actual['datos'].get(c),'despues':v} for c,v in cambios.items()]
                legados.append({'producto_id':actual['id'],'datos':cambios})
                f['accion']='actualizar' if cambios else 'sin_cambios'
            else:
                for h in tablas:tablas[h].extend(t[h])
                f['accion']='pendiente'
            if precio:precios.append(dict(precio,producto_id=actual['id'],sku=sku))
        except LimiteCoincidencias:raise
        except (ValueError,TypeError) as e:f['errores']=[str(e)]
        entrada_coincidente.append(r)
        vista.append(f)
    repetidos=Counter(f['sku'] for f in vista)
    for f in vista:
        if repetidos[f['sku']]>1:f['errores'].append('El mismo producto aparece más de una vez. Conservá una sola fila.');f['accion']='rechazado'
    vistas=completo.previsualizar(tablas,Producto=Producto,session=session,organizacion_id=organizacion_id) if tablas['Productos'] else []
    por_sku={f['sku']:f for f in vistas}
    for f in vista:
        if f['accion']=='pendiente':
            p=por_sku[f['sku']];f.update(accion=p['accion'],errores=p['errores'])
            f['cambios']+=p['detalle_cambios']
            f['cambios'] += [{'campo':c,'antes':'Lista actual','despues':'Lista de la planilla'} for c in p['cambios'] if c in ('Barras','Ubicaciones')]
            if any(p['sku']==f['sku'] and p['cambia'] for p in precios) and not f['errores']:f['accion']='actualizar'
    insumos=Counter(p['insumo_id'] for p in precios if p['insumo_id'])
    for f in vista:
        if any(p['sku']==f['sku'] and p['insumo_id'] and insumos[p['insumo_id']]>1 for p in precios):
            f['errores'].append('Varias filas actualizan la misma fuente de insumo. Conservá una sola.');f['accion']='rechazado'
    return {'filas':vista,'planilla':vistas,'legados':legados,'precios':precios,
            'entrada':entrada_coincidente,'campos':list(campos),'proveedor_id':proveedor_id,
            'leidas':len(filas),'ignoradas':ignoradas,'muestras_ignoradas':muestras_ignoradas}


def firmar(vista,org,usuario,secreto):
    return URLSafeTimedSerializer(secreto,salt=SALT).dumps({'org':org,'usuario':usuario,'vista':vista})


def aplicar(token,*,Producto,session,organizacion_id,usuario,secreto):
    if len(token)>1000000:raise ValueError('La revisión supera el tamaño permitido.')
    try:lote=URLSafeTimedSerializer(secreto,salt=SALT).loads(token,max_age=900)
    except BadSignature:raise ValueError('La revisión venció o cambió. Volvé a subir la planilla.') from None
    if lote['org']!=organizacion_id or lote['usuario']!=usuario.id:raise ValueError('La revisión pertenece a otro usuario u organización.')
    anterior=lote['vista']
    if not anterior['filas'] or any(f['errores'] for f in anterior['filas']):raise ValueError('Corregí los errores antes de confirmar.')
    try:
        ids=[f['producto_id'] for f in anterior['filas']]
        session.query(Producto).filter(Producto.organizacion_id==organizacion_id,Producto.id.in_(ids)).order_by(Producto.id).with_for_update().populate_existing().all()
        from models.producto_ficha_maestra import ProductoFichaMaestra
        session.query(ProductoFichaMaestra).filter(ProductoFichaMaestra.organizacion_id==organizacion_id,ProductoFichaMaestra.producto_id.in_(ids)).order_by(ProductoFichaMaestra.producto_id).with_for_update().populate_existing().all()
        from models.fuentes_costo_productivo import InsumoProductivo
        ids_insumo=sorted({p['insumo_id'] for p in anterior['precios'] if p['insumo_id']})
        if ids_insumo:
            session.query(InsumoProductivo).filter(InsumoProductivo.organizacion_id==organizacion_id,InsumoProductivo.id.in_(ids_insumo)).order_by(InsumoProductivo.id).with_for_update().populate_existing().all()
            from models.fuentes_costo_productivo import InsumoPrecioVersion
            session.query(InsumoPrecioVersion).filter(InsumoPrecioVersion.insumo_id.in_(ids_insumo)).order_by(InsumoPrecioVersion.id).with_for_update().populate_existing().all()
            from services import insumos_presentacion as presentacion
            if presentacion.disponible(session):
                from models.insumo_presentacion_compra import InsumoPresentacionCompra, InsumoPrecioBultoVersion
                session.query(InsumoPresentacionCompra).filter(InsumoPresentacionCompra.organizacion_id==organizacion_id,InsumoPresentacionCompra.insumo_id.in_(ids_insumo)).order_by(InsumoPresentacionCompra.insumo_id).with_for_update().populate_existing().all()
                session.query(InsumoPrecioBultoVersion).filter(InsumoPrecioBultoVersion.organizacion_id==organizacion_id,InsumoPrecioBultoVersion.precio_version_id.in_(session.query(InsumoPrecioVersion.id).filter(InsumoPrecioVersion.insumo_id.in_(ids_insumo)))).order_by(InsumoPrecioBultoVersion.precio_version_id).with_for_update().populate_existing().all()
        nuevo=previsualizar(anterior['entrada'],anterior['campos'],proveedor_id=anterior['proveedor_id'],Producto=Producto,session=session,organizacion_id=organizacion_id)
        if any(f['errores'] for f in nuevo['filas']) or [(f['producto_id'],f['huella']) for f in nuevo['filas']]!=[(f['producto_id'],f['huella']) for f in anterior['filas']]:
            raise ValueError('Cambió una ficha desde la revisión. Volvé a subir la planilla; no se guardó nada.')
        if completo._firma(nuevo['precios'])!=completo._firma(anterior['precios']):raise ValueError('Cambió el precio o bulto desde la revisión. Volvé a revisar.')
        if nuevo['planilla']:
            completo.aplicar(completo.firmar(nuevo['planilla'],organizacion_id,usuario.id,secreto),
                Producto=Producto,session=session,organizacion_id=organizacion_id,usuario=usuario,secreto=secreto,commit=False)
        for f in nuevo['legados']:
            p=session.get(Producto,f['producto_id'])
            for c,v in f['datos'].items():setattr(p,c,v)
        from services import insumos_presentacion as bultos
        from models.fuentes_costo_productivo import InsumoProductivo, InsumoPrecioVersion
        from services.fuentes_costo_productivo import registrar_precio_insumo
        from models.compras import ProveedorCompra
        for f in nuevo['precios']:
            if not f['insumo_id'] or not f['cambia']:continue
            insumo=session.query(InsumoProductivo).filter_by(id=f['insumo_id'],organizacion_id=organizacion_id).one()
            proveedor=session.get(ProveedorCompra,f['proveedor_id'])
            bultos.guardar(insumo,f['datos'],session,organizacion_id)
            version=registrar_precio_insumo(insumo,moneda=f['moneda'],precio_unitario_centavos=f['plan']['unitario_centavos'],
                proveedor_referencia=proveedor.razon_social,observacion='Lista externa desde Productos. Sin IVA.',
                creado_por_usuario_id=getattr(usuario,'id',None),InsumoPrecioVersion=InsumoPrecioVersion,db_session=session,commit=False)
            bultos.guardar_precio(version,insumo,f['datos'],session,organizacion_id)
        session.commit()
        return sum(f['accion']=='actualizar' for f in nuevo['filas'])
    except Exception:session.rollback();raise
