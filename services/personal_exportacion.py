"""Una descarga Excel de los legajos de la organización, sin secretos de acceso."""
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo
from sqlalchemy import inspect
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from models.personal_administrativo import PersonalDetalle, HorarioPersonal, AusenciaPersonal, PersonalAcceso
from services.personal_legajo import DIAS


def exportar_personal(organizacion, *, modelos, db_session, unidad_negocio_id=None):
    inspector=inspect(db_session.connection());pendientes=[]
    def disponible(modelo, nombre):
        ok=inspector.has_table(modelo.__tablename__)
        if not ok:pendientes.append(nombre)
        return ok
    E=modelos['EmpleadoProductivo']
    from services.personal_unidad import consulta
    q=consulta(organizacion.id,unidad_negocio_id,modelos,db_session) if unidad_negocio_id is not None else db_session.query(E).filter_by(organizacion_id=organizacion.id,tipo_registro='empleado')
    empleados=q.order_by(E.codigo,E.id).all()
    por_id={e.id:e for e in empleados};ids=list(por_id)
    unidades={u.id:u.nombre for u in db_session.query(modelos['UnidadNegocio']).filter_by(organizacion_id=organizacion.id).all()}
    wb=Workbook();wb.remove(wb.active)
    def hoja(nombre,titulos,filas,monetarias=()):
        ws=wb.create_sheet(nombre);ws.append(titulos)
        for valores in filas:
            ws.append([ILLEGAL_CHARACTERS_RE.sub('',v) if isinstance(v,str) else v for v in valores])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
                elif isinstance(cell.value,datetime):cell.number_format='dd/mm/yyyy hh:mm'
                elif isinstance(cell.value,date):cell.number_format='dd/mm/yyyy'
                elif isinstance(cell.value,time):cell.number_format='hh:mm'
                elif isinstance(cell.value,Decimal):cell.number_format='0.####'
                if cell.column in monetarias:cell.number_format='#,##0.00'
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='124E64');cell.alignment=Alignment(wrap_text=True)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.row_dimensions[1].height=30
        for columna in ws.columns:
            longitud=max((len(str(c.value or '')) for c in columna),default=10)
            ws.column_dimensions[columna[0].column_letter].width=min(45,max(14,longitud+2))
        return ws
    def identidad(empleado_id):
        e=por_id[empleado_id];return [e.codigo,e.nombre]
    def registros(nombre,clave):
        modelo=modelos[clave]
        if not disponible(modelo,nombre) or not ids:return []
        q=db_session.query(modelo).filter(modelo.empleado_id.in_(ids))
        if unidad_negocio_id is not None and hasattr(modelo,'unidad_negocio_id'):q=q.filter(modelo.unidad_negocio_id==unidad_negocio_id)
        return q.order_by(modelo.empleado_id,modelo.id).all()
    detalles={d.empleado_id:d for d in db_session.query(PersonalDetalle).filter(PersonalDetalle.empleado_id.in_(ids),PersonalDetalle.organizacion_id==organizacion.id).all()} if disponible(PersonalDetalle,'Datos ampliados') and ids else {}
    campos=['dni','cuil','fecha_nacimiento','fecha_ingreso','fecha_egreso','telefono','email','domicilio','localidad','provincia','contacto_emergencia','categoria','convenio','modalidad']
    titulos=['Legajo','Nombre y apellido','Estado','Unidad','Sector','Puesto','DNI','CUIL','Nacimiento','Ingreso','Egreso','Teléfono','Email','Domicilio','Localidad','Provincia','Contacto emergencia','Categoría','Convenio','Modalidad','Observación general']
    from services.personal_areas import unidades_asignadas
    enlaces_disponibles=inspector.has_table(modelos['RecursoEmpleadoProductivo'].__tablename__)
    unidades_personal=unidades_asignadas(empleados) if enlaces_disponibles else {e.id:unidades.get(e.unidad_negocio_id,'Compartido') for e in empleados}
    if unidad_negocio_id is not None:unidades_personal={e.id:unidades.get(unidad_negocio_id,'') for e in empleados}
    hoja('Personal',titulos,([e.codigo,e.nombre,'Activo' if e.activo else 'Inactivo',unidades_personal[e.id],e.sector,e.puesto]+[getattr(detalles.get(e.id),campo,None) for campo in campos]+[e.observacion] for e in empleados))
    dinero=lambda valor:Decimal(valor)/100
    filas=[]
    for v in registros('Sueldos','EmpleadoCostoVersion'):
        filas.append(identidad(v.empleado_id)+[v.numero_version,v.moneda,'Vigente' if v.vigente else 'Histórica',v.vigente_desde,v.vigente_hasta,
            dinero(v.sueldo_base_centavos),dinero(v.cargas_sociales_centavos),v.porcentaje_cargas,'General' if v.usa_porcentaje_general else 'Individual',
            dinero(v.adicionales_centavos),dinero(v.otros_costos_centavos),v.horas_mensuales,v.horas_productivas,v.porcentaje_productivo,
            dinero(v.costo_mensual_total_centavos),dinero(v.costo_hora_productiva_centavos),dinero(v.costo_minuto_productivo_centavos),v.ubicacion_trabajo,v.tipo_funcion,v.observacion,v.fecha_creacion])
    hoja('Sueldos',['Legajo','Empleado','Versión','Moneda','Estado','Desde','Hasta','Sueldo base','Cargas','Cargas %','Tipo de cargas','Adicionales','Otros costos','Horas mensuales','Horas para costo','Participación %','Costo mensual','Costo por hora','Costo por minuto','Ubicación','Función','Observación','Fecha de registro'],filas,monetarias=(8,9,12,13,17,18,19))
    horarios=db_session.query(HorarioPersonal).filter(HorarioPersonal.empleado_id.in_(ids)).order_by(HorarioPersonal.empleado_id,HorarioPersonal.dia_semana,HorarioPersonal.entrada,HorarioPersonal.id).all() if disponible(HorarioPersonal,'Horarios') and ids else []
    hoja('Horarios',['Legajo','Empleado','Día','Entrada','Salida','Salida al día siguiente','Descanso (min)','Desde','Hasta','Estado','Observación','Fecha de registro'],(identidad(h.empleado_id)+[DIAS[h.dia_semana],h.entrada,h.salida,h.nocturno,h.pausa_minutos,h.desde,h.hasta,h.estado,h.observacion,h.fecha_creacion] for h in horarios))
    novedades=registros('Novedades','NovedadLaboralPreparatoria');novedades=[n for n in novedades if n.organizacion_id==organizacion.id];novedades_por_id={n.id:n for n in novedades}
    ausencias=db_session.query(AusenciaPersonal).filter(AusenciaPersonal.empleado_id.in_(ids)).order_by(AusenciaPersonal.empleado_id,AusenciaPersonal.id).all() if disponible(AusenciaPersonal,'Vacaciones y licencias') and ids else []
    filas=[];vinculadas=set()
    for a in ausencias:
        n=novedades_por_id.get(a.novedad_id)
        if n is None or n.empleado_id!=a.empleado_id:continue
        vinculadas.add(n.id);filas.append(identidad(a.empleado_id)+[a.tipo,n.fecha_desde,n.fecha_hasta,a.periodo,a.dias_computados,n.detalle,n.estado,unidades.get(n.unidad_negocio_id,'')])
    for n in novedades:
        if n.id not in vinculadas and n.tipo in {'ausencia','licencia'}:filas.append(identidad(n.empleado_id)+[n.tipo,n.fecha_desde,n.fecha_hasta,None,None,n.detalle,n.estado,unidades.get(n.unidad_negocio_id,'')])
    hoja('Vacaciones y licencias',['Legajo','Empleado','Tipo','Desde','Hasta','Período','Días computados','Detalle','Estado','Unidad'],filas)
    hoja('Novedades',['Legajo','Empleado','Tipo','Desde','Hasta','Detalle','Minutos extra propuestos','Estado','Aprobada','Afecta nómina','Unidad','Observación','Fecha de registro'],(identidad(n.empleado_id)+[n.tipo,n.fecha_desde,n.fecha_hasta,n.detalle,n.minutos_propuestos,n.estado,n.aprobada,n.afecta_nomina,unidades.get(n.unidad_negocio_id,''),n.observacion,n.fecha_creacion] for n in novedades))
    turnos=[t for t in registros('Turnos','TurnoLaboralPlanificado') if t.organizacion_id==organizacion.id]
    hoja('Turnos',['Legajo','Empleado','Inicio','Fin','Sector','Tarea','Estado','Confirmado','Afecta nómina','Unidad','Observación','Fecha de registro'],(identidad(t.empleado_id)+[t.inicio,t.fin,t.sector,t.tarea,t.estado,t.confirmado,t.afecta_nomina,unidades.get(t.unidad_negocio_id,''),t.observacion,t.fecha_creacion] for t in turnos))
    competencias=[c for c in registros('Competencias','CompetenciaEmpleadoPreparatoria') if c.organizacion_id==organizacion.id]
    hoja('Competencias',['Legajo','Empleado','Código','Competencia o certificación','Nivel','Vencimiento','Unidad de referencia','Verificada','Habilita asignación','Observación','Fecha de registro'],(identidad(c.empleado_id)+[c.codigo,c.nombre,c.nivel,c.vence_el,unidades.get(c.unidad_negocio_id,''),c.verificada,c.habilita_asignacion,c.observacion,c.fecha_creacion] for c in competencias))
    distribuciones=[d for d in registros('Distribución de costos','EmpleadoDistribucionVersion') if d.organizacion_id==organizacion.id]
    hoja('Distribución de costos',['Legajo','Empleado','Revisión','Unidad','Asignación %','Ubicación','Función','Vigente','Desde','Hasta','Observación'],(identidad(d.empleado_id)+[d.numero_revision,unidades.get(d.unidad_negocio_id,''),d.porcentaje_asignacion,d.ubicacion_trabajo,d.tipo_funcion,d.vigente,d.vigente_desde,d.vigente_hasta,d.observacion] for d in distribuciones))
    vinculos=registros('Equipos','RecursoEmpleadoProductivo')
    recursos={r.id:r for r in db_session.query(E).filter(E.organizacion_id==organizacion.id,E.tipo_registro=='recurso',E.id.in_({v.recurso_id for v in vinculos})).all() if unidad_negocio_id is None or r.unidad_negocio_id==unidad_negocio_id} if vinculos else {}
    hoja('Equipos',['Legajo','Empleado','Código equipo','Equipo','Dedicación %','Observación','Unidad del área','Sector'],(identidad(v.empleado_id)+[recursos[v.recurso_id].codigo,recursos[v.recurso_id].nombre,v.porcentaje_dedicacion,v.observacion,unidades.get(recursos[v.recurso_id].unidad_negocio_id,''),recursos[v.recurso_id].sector] for v in vinculos if v.recurso_id in recursos))
    from services.personal_legajo import version_actual
    recursos_todos=db_session.query(E).filter_by(organizacion_id=organizacion.id,tipo_registro='recurso').order_by(E.unidad_negocio_id,E.sector).all()
    if unidad_negocio_id is not None:recursos_todos=[r for r in recursos_todos if r.unidad_negocio_id==unidad_negocio_id]
    filas_areas=[]
    for r in recursos_todos:
        v=version_actual(r) if inspector.has_table(modelos['EmpleadoCostoVersion'].__tablename__) else None
        filas_areas.append([unidades.get(r.unidad_negocio_id,''),r.codigo,r.nombre,r.sector,len(r.miembros_recurso) if enlaces_disponibles else None,
            dinero(v.costo_mensual_total_centavos) if v else None,v.horas_productivas if v else None,
            dinero(v.costo_hora_productiva_centavos) if v else None,dinero(v.costo_minuto_productivo_centavos) if v else None,
            'Activo' if r.activo else 'Inactivo',r.porcentaje_indirecto])
    hoja('Áreas mano de obra',['Unidad','Código','Área','Sector','Integrantes','Costo mensual','Horas asignadas','Costo por hora','Costo por minuto','Estado','Recargo adicional %'],filas_areas,monetarias=(6,8,9))
    accesos={a.empleado_id:a for a in db_session.query(PersonalAcceso).filter(PersonalAcceso.empleado_id.in_(ids)).all()} if disponible(PersonalAcceso,'Accesos') and ids else {}
    M=modelos['UsuarioOrganizacion'];U=modelos['UsuarioSistema'];cuentas={}
    if disponible(M,'Usuarios de la organización') and disponible(U,'Usuarios del sistema'):
        cuentas={mem.id:(mem,user) for mem,user in db_session.query(M,U).join(U,U.id==M.usuario_id).filter(M.organizacion_id==organizacion.id,M.id.in_([a.usuario_organizacion_id for a in accesos.values() if a.usuario_organizacion_id])).all()}
    filas=[]
    for e in empleados:
        a=accesos.get(e.id);par=cuentas.get(a.usuario_organizacion_id) if a else None
        mem,user=par if par else (None,None)
        filas.append(identidad(e.id)+[user.username if user else None,mem.rol if mem else None,bool(user.activo and mem.activa) if user else False,bool(a and a.codigo_alarma_cifrado)])
    hoja('Accesos',['Legajo','Empleado','Usuario Sistema Fierro','Rol actual','Acceso habilitado','Código de alarma configurado'],filas)
    ahora=datetime.now(ZoneInfo('America/Argentina/Buenos_Aires')).replace(tzinfo=None)
    info=hoja('Información',['Dato','Detalle'],[
        ['Organización',organizacion.nombre],['Fecha de descarga (Argentina)',ahora],['Cantidad de empleados',len(empleados)],
        ['Alcance',('Unidad activa: '+unidades.get(unidad_negocio_id,str(unidad_negocio_id))+'; legajos vinculados, activos e inactivos. Ficha, sueldo y horario comunes de cada legajo; operaciones y distribuciones sólo de esta unidad.') if unidad_negocio_id is not None else 'Toda la organización, todas las unidades, activos e inactivos; incluye historial.'],
        ['Accesos','No se exportan contraseñas, hashes ni códigos de alarma. Sólo usuario, rol, estado y si la alarma está configurada.'],
        ['Fechas','Horarios y fechas laborales son locales; fechas de registro del sistema y versiones se muestran en UTC.'],
        ['Fuentes sin preparar',', '.join(pendientes) if pendientes else 'Todas disponibles']])
    wb.active=wb.index(info)
    salida=BytesIO();wb.save(salida);salida.seek(0)
    return salida
