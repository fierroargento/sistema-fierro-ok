"""Contrato XLSX Fierro para controlar y editar productos existentes.

Exportación e importación comparten columnas y versión. No ejecuta fórmulas,
no crea productos ni borra fichas por omisiones. Imágenes e historial se conservan.
"""
from collections import defaultdict
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
import json
from zipfile import ZipFile, BadZipFile
from itsdangerous import URLSafeTimedSerializer, BadSignature

from services.productos_ficha_maestra import construir_filas, normalizar_ficha, validar_vinculos
from services.productos_datos_generales import leer_formulario_completo

FORMATO='FIERRO-PRODUCTOS-1'
BASE=('sku','nombre','linea','descripcion','tipo','unidad_medida','activo','no_stockeable','permitir_stock_negativo','habilita_compra','habilita_venta','habilita_insumo','habilita_componente','insumo_id','peso_gr','alto_cm','ancho_cm','largo_cm','permite_correo','permite_via_cargo','requiere_revision_logistica','observacion_logistica')
GENERAL=('rubro','subrubro','marca','moneda_costo','precio_costo','iva_pct','iibb_pct','utilidad_pct','utilidad_base','sucursales_definidas','depositos_definidos')
COLUMNAS={
 'Productos':('producto_id','huella')+BASE+GENERAL,
 'Perfiles':('sku','perfil_id','tipo','unidad_id','activo','modo','tarifa_id'),
 'Componentes':('sku','perfil_id','nombre','insumo_id','cantidad','modo','largo_cm','rendimiento'),
 'Tareas':('sku','perfil_id','nombre','segundos'),
 'Consumos':('sku','perfil_id','id','insumo_id','cantidad','porcentaje_merma','observacion'),
 'Operaciones':('sku','perfil_id','id','empleado_id','nombre','minutos','observacion'),
 'Maquinas':('sku','perfil_id','id','maquina_id','nombre','minutos','observacion'),
 'Indirectos':('sku','perfil_id','id','costo_fijo_id','porcentaje_asignacion','unidades_mensuales','observacion'),
 'Combos':('sku','perfil_id','componente_sku','cantidad'),
 'Proveedores':('sku','proveedor_id','codigo'),
 'Ubicaciones':('sku','tipo','ubicacion_id'),
 'Barras':('sku','codigo'),
}
BOOLEANOS={'activo','no_stockeable','permitir_stock_negativo','habilita_compra','habilita_venta','habilita_insumo','habilita_componente','permite_correo','permite_via_cargo','requiere_revision_logistica','sucursales_definidas','depositos_definidos'}
NUMEROS={'peso_gr','alto_cm','ancho_cm','largo_cm','precio_costo','iva_pct','iibb_pct','utilidad_pct','cantidad','largo_cm','rendimiento','segundos','minutos','porcentaje_merma','porcentaje_asignacion','unidades_mensuales'}

ETIQUETAS={'sku':'SKU','producto_id':'ID interno','huella':'Control de versión','nombre':'Nombre','linea':'Línea','descripcion':'Descripción',
 'tipo':'Tipo','unidad_medida':'Unidad de medida','activo':'Activo','no_stockeable':'No stockeable','permitir_stock_negativo':'Permitir stock negativo',
 'habilita_compra':'Disponible para compra','habilita_venta':'Disponible para venta','habilita_insumo':'Como insumo','habilita_componente':'Como componente',
 'insumo_id':'Insumo ID','precio_costo':'Costo simple sin IVA','moneda_costo':'Moneda del costo','iva_pct':'IVA (%)','iibb_pct':'IIBB (%)','utilidad_pct':'Utilidad (%)',
 'utilidad_base':'Utilidad sobre','perfil_id':'Perfil ID','unidad_id':'Unidad ID','tarifa_id':'Tarifa ID','componente_sku':'SKU componente','proveedor_id':'Proveedor ID',
 'codigo':'Código','ubicacion_id':'Ubicación ID','largo_cm':'Largo (cm)','segundos':'Segundos por unidad','minutos':'Minutos por unidad','rendimiento':'Piezas por unidad de insumo',
 'porcentaje_merma':'Merma (%)','porcentaje_asignacion':'Asignación (%)','unidades_mensuales':'Unidades mensuales','peso_gr':'Peso (g)','alto_cm':'Alto (cm)','ancho_cm':'Ancho (cm)',
 'permite_correo':'Permite Correo','permite_via_cargo':'Permite Vía Cargo','requiere_revision_logistica':'Revisar logística','observacion_logistica':'Observación logística'}

def etiqueta(c):return ETIQUETAS.get(c,c.replace('_',' ').capitalize())


def _normal(v):
    if isinstance(v,Decimal):return format(v.normalize(),'f')
    if isinstance(v,dict):return {k:_normal(x) for k,x in v.items()}
    if isinstance(v,list):return [_normal(x) for x in v]
    return v


def _firma(obj):
    return json.dumps(_normal(obj),sort_keys=True,ensure_ascii=False,separators=(',',':'))


def actuales(Producto,session,org):
    productos=session.query(Producto).filter_by(organizacion_id=org).order_by(Producto.sku).all()
    disponible,filas=construir_filas(productos,session,org)
    if not disponible:raise ValueError('Falta preparar el maestro de productos.')
    salida={}
    for f in filas:
        sku=f['datos']['sku'].strip().upper()
        if sku in salida:raise ValueError(f'{sku}: hay identidades repetidas en el maestro; revisalas antes de exportar.')
        f['datos']['sku']=sku;salida[sku]=f
    return salida


def tablas_de_filas(filas):
    tablas={k:[] for k in COLUMNAS}
    def agregar(hoja,**datos):tablas[hoja].append({c:datos.get(c,'') for c in COLUMNAS[hoja]})
    for f in filas.values():
        d=f['datos'];a=f['adicionales'];sku=d['sku']
        agregar('Productos',producto_id=f['id'],huella=f['huella'],**{c:d.get(c,'') for c in BASE},**{c:a.get(c,'ARS' if c=='moneda_costo' else 'costo' if c=='utilidad_base' else False if c in ('sucursales_definidas','depositos_definidos') else '') for c in GENERAL})
        for p in f['producciones']:
            des=p.get('despiece');fisico=bool(des and des['actual'])
            if des and not des['actual']:raise ValueError(f'{sku}: el despiece está desactualizado. Revisalo antes de exportar para editar.')
            agregar('Perfiles',sku=sku,perfil_id=p['perfil_id'],tipo='produccion',unidad_id=p['unidad_id'],activo=p['activo'],modo='piezas' if fisico else 'consumos',tarifa_id=des['datos'].get('tarifa_id') if fisico else '')
            if fisico:
                for pieza in des['datos']['piezas']:agregar('Componentes',sku=sku,perfil_id=p['perfil_id'],**pieza)
                for tarea in des['datos']['tareas']:agregar('Tareas',sku=sku,perfil_id=p['perfil_id'],**tarea)
            for grupo,hoja in [('insumos','Consumos'),('operaciones','Operaciones'),('maquinas','Maquinas'),('fijos','Indirectos')]:
                if fisico and grupo in ('insumos','operaciones'):continue
                for linea in p[grupo]:agregar(hoja,sku=sku,perfil_id=p['perfil_id'],**{k:v for k,v in linea.items() if k in COLUMNAS[hoja]})
        for p in f['composiciones']:
            agregar('Perfiles',sku=sku,perfil_id=p['perfil_id'],tipo='combo',unidad_id=p['unidad_combo_id'],activo=p['activo'],modo='combo')
            for c in p['componentes_combo']:agregar('Combos',sku=sku,perfil_id=p['perfil_id'],componente_sku=c['sku'],cantidad=c['cantidad'])
        for p in a.get('proveedores',[]):agregar('Proveedores',sku=sku,proveedor_id=p['proveedor_id'],codigo=p['codigo'])
        for tipo in ('sucursales','depositos'):
            for ident in a.get(tipo,[]):agregar('Ubicaciones',sku=sku,tipo='sucursal' if tipo=='sucursales' else 'deposito',ubicacion_id=ident)
        for codigo in a.get('codigos_barras',[]):agregar('Barras',sku=sku,codigo=codigo)
    return tablas


def exportar(*,Producto,session,organizacion_id,base_url=""):
    from services.productos_datos_generales import esquema_datos_disponible
    if not esquema_datos_disponible(session):raise ValueError('Falta preparar la ficha completa para exportar e importar la planilla Fierro.')
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment,Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    font_editable=Font(color='165DA8',size=11);font_lectura=Font(color='64748B',size=11)
    fondo_alterno=PatternFill('solid',fgColor='F0F5F8');alineacion=Alignment(vertical='top',wrap_text=True)
    protegido=Protection(locked=True);editable=Protection(locked=False)
    filas=actuales(Producto,session,organizacion_id)
    if len(filas)>500:raise ValueError('Esta exportación admite hasta 500 productos por archivo.')
    tablas=tablas_de_filas(filas);libro=Workbook();libro.remove(libro.active)
    control=libro.create_sheet('Control');control.append(['formato','organizacion_id','fecha_utc']);control.append([FORMATO,organizacion_id,datetime.now(timezone.utc).isoformat()])
    for nombre,columnas in COLUMNAS.items():
        ws=libro.create_sheet(nombre);ws.append([etiqueta(c) for c in columnas])
        for fila in tablas[nombre]:
            valores=[]
            for c in columnas:
                v=fila[c]
                if c in BOOLEANOS:v='Sí' if v else 'No'
                elif c in NUMEROS and v not in (None,''):
                    numero=Decimal(str(v));v=float(numero) if len(numero.as_tuple().digits)<=15 else format(numero,'f')
                valores.append(v)
            ws.append(valores)
            for n,c in enumerate(columnas,1):
                celda=ws.cell(ws.max_row,n)
                if isinstance(valores[n-1],str):celda.data_type='s' # nunca exportar texto como fórmula
                celda.protection=protegido if c in ('producto_id','huella','sku','perfil_id','id') or c=='tipo' and nombre=='Productos' or nombre=='Perfiles' and c!='tarifa_id' or c=='insumo_id' and nombre=='Productos' else editable
                celda.number_format='0.######' if c in NUMEROS and not isinstance(valores[n-1],str) else '@'
        for n,c in enumerate(columnas,1):
            letra=ws.cell(1,n).column_letter
            ws.column_dimensions[letra].width=36 if c in ('nombre','descripcion','observacion','observacion_logistica') else 22 if c in GENERAL or c in BOOLEANOS else 17
            if c in ('producto_id','huella'):ws.column_dimensions[letra].hidden=True
            if c in BOOLEANOS:
                dv=DataValidation(type='list',formula1='"Sí,No"');dv.errorTitle='Elegí Sí o No';dv.error='Usá una de las opciones de la lista.';dv.showErrorMessage=True;dv.errorStyle='stop';ws.add_data_validation(dv);dv.add(f'{letra}2:{letra}5001')
    referencias=libro.create_sheet('Referencias');referencias.append(['tipo','id','codigo','nombre','unidad'])
    from services.productos_produccion_maestro import modelos,disponible
    from models.unidad_negocio import UnidadNegocio
    from models.compras import ProveedorCompra
    from models.sucursal_operativa import SucursalOperativa
    for tipo,Modelo in [('unidad',UnidadNegocio),('proveedor',ProveedorCompra),('ubicacion',SucursalOperativa)]+([(g,m[1]) for g,m in modelos()[1].items()] if disponible(session) else []):
        for r in session.query(Modelo).filter_by(organizacion_id=organizacion_id):referencias.append([tipo,r.id,getattr(r,'codigo',''),getattr(r,'nombre',getattr(r,'razon_social','')),getattr(r,'unidad_medida','')])
    costos=libro.create_sheet('Costos de referencia');costos.append(['SKU','Unidad ID','Moneda','Versión','Tipo','Estado','Vigente','Costo sin IVA'])
    from models.costo_producto_version import CostoProductoVersion
    por_id={f['id']:sku for sku,f in filas.items()}
    for r in session.query(CostoProductoVersion).filter(CostoProductoVersion.organizacion_id==organizacion_id,CostoProductoVersion.producto_id.in_(por_id)).order_by(CostoProductoVersion.producto_id,CostoProductoVersion.numero_version):
        costos.append([por_id[r.producto_id],r.unidad_negocio_id,r.moneda,r.numero_version,r.tipo,r.estado,'Sí' if r.vigente else 'No',r.costo_total_centavos/100])
    imagenes=libro.create_sheet('Imágenes de consulta');imagenes.append(['SKU','Nombre de archivo','Principal','Ver en Fierro'])
    for sku,f in filas.items():
        for r in f['adicionales'].get('imagenes',[]):imagenes.append([sku,r['nombre'],'Sí' if r['principal'] else 'No',base_url.rstrip('/')+r['url']])
    for ws in (control,referencias,costos,imagenes):ws.protection.sheet=True
    ayuda=libro.create_sheet('Ayuda');ayuda.append(['Uso de la planilla Fierro'])
    for t in ['Editá las celdas azules y guardá como XLSX. Luego usá Importar planilla Fierro en Productos.','No cambies SKU, producto_id, huella ni los datos de identidad de Perfiles. Son productos existentes, no altas.','Para quitar un producto de uso cambiá Activo a No. Quitar su fila del archivo no borra el producto.','Para fabricación con piezas editá Componentes y Tareas; el consumo y costo se recalculan en Fierro. Para fichas anteriores sin piezas, usá Consumos y Operaciones.','Perfiles identifica la unidad de cada receta; no elimines perfiles. Referencias contiene los códigos e identificadores para elegir los recursos.','Precio de costo es editable sólo para productos simples. Para producción y combos se construye desde su composición. Importes sin IVA; IVA, IIBB y utilidad se expresan en puntos porcentuales: 21, 3.5, 90.','Proveedores, Ubicaciones y Barras reemplazan esas listas sólo de los SKU incluidos. Para quitar una relación eliminá su fila; deben conservarse todas las hojas.','Imágenes, existencias, precios publicados e historial se conservan en el sistema. No se modifican a través de esta planilla.','La importación rechaza fichas modificadas en el sistema desde esta exportación. Descargá de nuevo si quedó desactualizada.','Costos de referencia muestra el historial guardado, sin IVA y por unidad. Es de consulta y no modifica costos históricos ni precios publicados.','Los costos de los recursos se consultan en Fierro; Referencias es de consulta y no actualiza tarifas ni precios de insumos.']:
        ayuda.append([t])
    ayuda.column_dimensions['A'].width=110
    referencias.column_dimensions['D'].width=48
    for ws in libro:
        if ws.title not in COLUMNAS and ws.title!='Ayuda':
            for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=24
        ws.freeze_panes='C2' if ws.title=='Productos' else 'B2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17566D');cell.alignment=Alignment(vertical='center',wrap_text=True)
        ws.row_dimensions[1].height=32
        for fila in ws.iter_rows(min_row=2):
            for cell in fila:
                if isinstance(cell.value,str):cell.data_type='s'
                if cell.row%2==0:cell.fill=fondo_alterno
                cell.font=font_lectura if cell.protection.locked else font_editable;cell.alignment=alineacion
    salida=BytesIO();libro.save(salida);salida.seek(0);return salida


def leer(archivo,org):
    from openpyxl import load_workbook
    contenido=archivo.read(4*1024*1024+1)
    if len(contenido)>4*1024*1024 or not str(archivo.filename).lower().endswith('.xlsx'):raise ValueError('Elegí la planilla Fierro XLSX, de hasta 4 MB.')
    try:
        with ZipFile(BytesIO(contenido)) as z:
            if len(z.infolist())>2000 or sum(i.file_size for i in z.infolist())>30*1024*1024:raise ValueError('La planilla descomprimida supera el límite.')
        w=load_workbook(BytesIO(contenido),read_only=True,data_only=False,keep_links=False)
    except (BadZipFile,KeyError):raise ValueError('No se pudo leer la planilla XLSX.') from None
    try:
        if 'Control' not in w:raise ValueError('Usá Exportar para editar; esta no es una planilla Fierro.')
        meta=list(w['Control'].values)
        if len(meta)<2 or meta[1][0]!=FORMATO or str(meta[1][1])!=str(org):raise ValueError('La planilla no pertenece a esta organización o versión de Fierro.')
        tablas={}
        for hoja,columnas in COLUMNAS.items():
            if hoja not in w:raise ValueError(f'Falta la hoja {hoja}; no se aplicó ningún cambio.')
            ws=w[hoja]
            if ws.max_row>50001 or ws.max_column>80:raise ValueError(f'{hoja}: supera los límites permitidos.')
            fuente=ws.values;enc=next(fuente,())
            if tuple(enc)!=tuple(etiqueta(c) for c in columnas):raise ValueError(f'{hoja}: no cambies los encabezados ni su orden.')
            filas=[]
            for n,celdas in enumerate(ws.iter_rows(min_row=2),2):
                if any(c.data_type=='f' for c in celdas):raise ValueError(f'{hoja}, fila {n}: usá valores, no fórmulas.')
                r=[c.value for c in celdas]
                if not any(v not in (None,'') for v in r):continue
                if any(not isinstance(v,(str,int,float,bool,type(None))) for v in r):raise ValueError(f'{hoja}, fila {n}: usá valores, no fórmulas.')
                datos=dict(zip(columnas,(v if v is not None else '' for v in r)))
                for c in NUMEROS & datos.keys():
                    if isinstance(datos[c],(int,float)) and not isinstance(datos[c],bool):datos[c]=format(Decimal(str(datos[c])),'f')
                filas.append(datos)
            tablas[hoja]=filas
        if not tablas['Productos'] or len(tablas['Productos'])>500:raise ValueError('Incluí entre 1 y 500 productos existentes.')
        skus=[str(f['sku']).strip().upper() for f in tablas['Productos']]
        if len(skus)!=len(set(skus)):raise ValueError('Hay SKU repetidos en Productos.')
        for hoja,filas in tablas.items():
            for f in filas:f['sku']=str(f['sku']).strip().upper()
            if hoja!='Productos' and any(f['sku'] not in skus for f in filas):raise ValueError(f'{hoja}: hay filas de SKU que no están en Productos.')
        return tablas
    finally:w.close()


def _canon(tablas):
    from services.productos_ficha_maestra import booleano
    salida={}
    for hoja,filas in tablas.items():
        salida[hoja]=[]
        for f in filas:
            r={}
            for c,v in f.items():
                if c in BOOLEANOS:r[c]=booleano(v,c) if v not in (None,'') else ''
                elif c in NUMEROS and v not in (None,''):
                    try:
                        d=Decimal(str(v).replace(',','.'))
                        if not d.is_finite():raise ValueError()
                        r[c]=format(d.normalize(),'f')
                    except Exception:raise ValueError(f'{hoja}: {c} debe ser numérico.') from None
                else:r[c]='' if v is None else str(v)
            salida[hoja].append(r)
    return salida


def _por_sku(tablas,sku):return {h:[r for r in filas if r['sku']==sku] for h,filas in tablas.items()}


def _adicional(entrada,actual,tablas):
    from services.productos_ficha_maestra import booleano
    datos={c:entrada.get(c,'') for c in GENERAL}
    for c in ('sucursales_definidas','depositos_definidos'):datos[c]='1' if booleano(datos[c],c) else '0'
    datos.update(datos_generales_presentados='1',costo_maestro_presentado='1',iibb_maestro_presentado='1',
        proveedores_json=json.dumps([{k:r[k] for k in ('proveedor_id','codigo')} for r in tablas['Proveedores']]),
        sucursales_json=json.dumps([r['ubicacion_id'] for r in tablas['Ubicaciones'] if r['tipo']=='sucursal']),
        depositos_json=json.dumps([r['ubicacion_id'] for r in tablas['Ubicaciones'] if r['tipo']=='deposito']),
        codigos_barras_json=json.dumps([r['codigo'] for r in tablas['Barras']]),
        imagenes_conservar_json=json.dumps([r['id'] for r in actual['adicionales'].get('imagenes',[])]),
        imagen_principal=next((str(r['id']) for r in actual['adicionales'].get('imagenes',[]) if r['principal']),''))
    if any(r['tipo'] not in ('sucursal','deposito') for r in tablas['Ubicaciones']):raise ValueError('Ubicaciones: usá sucursal o deposito.')
    return datos


def _planes(tablas,actual,session,org,lectura):
    from services.productos_produccion_maestro import preparar,calcular
    from services.productos_combo_maestro import preparar_composicion
    from models.producto import Producto as ProductoModelo
    producciones=[];combos=[]
    perfiles={str(p['perfil_id']):p for p in tablas['Perfiles']}
    if len(perfiles)!=len(tablas['Perfiles']):raise ValueError('Perfiles: hay identificadores repetidos.')
    conocidos={str(p['perfil_id']):p for p in tablas_de_filas({actual['datos']['sku']:actual})['Perfiles']}
    if set(perfiles)!=set(conocidos):raise ValueError('No agregues ni elimines perfiles; conservá las unidades de fabricación y combos exportadas.')
    for ident,p in perfiles.items():
        original=conocidos[ident]
        if any(str(p[c] or '')!=str(original[c] or '') for c in ('sku','perfil_id','tipo','unidad_id','modo')):
            raise ValueError('No cambies la identidad, tipo, modo o unidad de un perfil exportado.')
        from services.productos_ficha_maestra import booleano
        if booleano(p['activo'],'activo')!=original['activo']:raise ValueError('La activación de perfiles se gestiona en Costos.')
        def lineas(hoja):return [{c:v for c,v in r.items() if c not in ('sku','perfil_id')} for r in tablas[hoja] if str(r['perfil_id'])==ident]
        if p['tipo']=='produccion':
            receta={'unidad_id':p['unidad_id'] or None,'maquinas':lineas('Maquinas'),'fijos':lineas('Indirectos')}
            if p['modo']=='piezas':
                if lineas('Consumos') or lineas('Operaciones'):raise ValueError('Las fichas con piezas usan Componentes y Tareas, sin consumos u operaciones duplicados.')
                receta.update(piezas=lineas('Componentes'),tareas=lineas('Tareas'),tarifa_id=p['tarifa_id'] or None)
            else:
                if lineas('Componentes') or lineas('Tareas'):raise ValueError('Esta ficha usa Consumos y Operaciones; no cambies su método de fabricación desde Excel.')
                receta.update(insumos=lineas('Consumos'),operaciones=lineas('Operaciones'))
            # Sólo se prepara/costea una receta si su sección realmente cambió.
            hojas=('Componentes','Tareas','Consumos','Operaciones','Maquinas','Indirectos')
            anteriores=tablas_de_filas({actual['datos']['sku']:actual})
            antes={h:[r for r in anteriores[h] if str(r['perfil_id'])==ident] for h in hojas}
            despues={h:[r for r in tablas[h] if str(r['perfil_id'])==ident] for h in hojas}
            if _canon(antes)==_canon(despues) and str(p['tarifa_id'] or '')==str(original['tarifa_id'] or ''):continue
            plan=preparar(receta,session,org,actual['id'],lectura=lectura);costo=calcular(plan,org)
            producciones.append({'perfil_id':int(ident),'receta':receta,'huella_costo':costo['huella'],'costo_centavos':costo['total_centavos']})
        else:
            anteriores=tablas_de_filas({actual['datos']['sku']:actual})
            if _canon({'Combos':[r for r in anteriores['Combos'] if str(r['perfil_id'])==ident]})==_canon({'Combos':[r for r in tablas['Combos'] if str(r['perfil_id'])==ident]}):continue
            filas=lineas('Combos');datos=dict(actual['datos'],componentes_combo=[{'sku':r['componente_sku'],'cantidad':r['cantidad']} for r in filas],unidad_combo_id=p['unidad_id'] or None)
            datos=normalizar_ficha(datos);preparar_composicion(session.get(ProductoModelo,actual['id']),datos,session,org)
            combos.append(datos)
    for hoja in ('Componentes','Tareas','Consumos','Operaciones','Maquinas','Indirectos','Combos'):
        if any(str(r['perfil_id']) not in perfiles for r in tablas[hoja]):raise ValueError(f'{hoja}: perfil desconocido.')
        if hoja=='Combos' and any(perfiles[str(r['perfil_id'])]['tipo']!='combo' for r in tablas[hoja]):raise ValueError('Combos: el perfil no es de combo.')
        if hoja!='Combos' and any(perfiles[str(r['perfil_id'])]['tipo']!='produccion' for r in tablas[hoja]):raise ValueError(f'{hoja}: el perfil no es de producción.')
    return producciones,combos


def previsualizar(tablas,*,Producto,session,organizacion_id):
    from services.productos_datos_generales import validar_datos_generales
    from services.lectura_fabricacion_lote import LecturaFabricacion
    actuales_por_sku=actuales(Producto,session,organizacion_id);lectura=LecturaFabricacion(session,organizacion_id)
    vista=[]
    for entrada in tablas['Productos']:
        sku=entrada['sku'];actual=actuales_por_sku.get(sku)
        fila={'sku':sku,'numero':len(vista)+2,'accion':'rechazado','errores':[],'cambios':[],'datos':None,'producto_id':None,'huella':entrada['huella'],'adicionales':None,'producciones':[],'combos':[],'detalle_cambios':[],'tablas':_por_sku(tablas,sku)}
        try:
            if actual is None or str(actual['id'])!=str(entrada['producto_id']):raise ValueError('El SKU o identificador no corresponde a un producto existente de esta organización.')
            if actual['huella']!=entrada['huella']:raise ValueError('El producto cambió desde la exportación; descargá una planilla nueva.')
            fila['producto_id']=actual['id']
            esperadas=tablas_de_filas({sku:actual})
            if _canon(esperadas)==_canon(fila['tablas']):fila['accion']='sin_cambios';vista.append(fila);continue
            if entrada['tipo']!=actual['datos']['tipo'] or str(entrada['insumo_id'] or '')!=str(actual['datos']['insumo_id'] or ''):raise ValueError('El tipo y vínculo de insumo se conservan; usá Gestionar para revisarlos.')
            datos=normalizar_ficha({c:entrada[c] for c in BASE})
            producto=session.get(Producto,actual['id'])
            adicional=_adicional(entrada,actual,fila['tablas']);parseado=leer_formulario_completo(adicional,[])
            if datos['tipo']!='simple' and (parseado['precio_costo_centavos'] is not None or parseado['moneda_costo']!='ARS'):raise ValueError('El costo de producción y combo se construye; no cargues precio manual ni otra moneda.')
            cambios=[];detalle_cambios=[]
            for c in BASE+GENERAL:
                if _canon({'x':[{c:entrada[c]}]})!=_canon({'x':[{c:esperadas['Productos'][0][c]}]}):
                    cambios.append(c.replace('_',' '));detalle_cambios.append({'campo':etiqueta(c),'antes':esperadas['Productos'][0][c],'despues':entrada[c]})
            for h in COLUMNAS:
                if h!='Productos' and _canon({h:esperadas[h]})!=_canon({h:fila['tablas'][h]}):cambios.append(h)
            rapidos={'nombre','linea','descripcion','rubro','subrubro','marca','iva_pct','iibb_pct','utilidad_pct','utilidad_base'}
            solo_atributos=bool(actual['adicionales'].get('revision')) and set(cambios)<={c.replace('_',' ') for c in rapidos}
            if not solo_atributos:
                validar_vinculos(producto,datos,session,organizacion_id)
                validar_datos_generales(producto,parseado,session)
            hojas_receta=('Perfiles','Componentes','Tareas','Consumos','Operaciones','Maquinas','Indirectos','Combos')
            receta_cambiada=any(h in cambios for h in hojas_receta)
            producciones,combos=_planes(fila['tablas'],actual,session,organizacion_id,lectura) if receta_cambiada else ([],[])
            fila.update(accion='actualizar',datos=datos,adicionales=adicional,producciones=producciones,combos=combos,cambios=cambios,detalle_cambios=detalle_cambios,solo_atributos=solo_atributos)
        except (ValueError,TypeError,KeyError) as error:fila['errores']=[str(error)]
        vista.append(fila)
    propietarios=defaultdict(set)
    for f in vista:
        for r in f['tablas']['Barras']:propietarios[('barra',str(r['codigo']))].add(f['sku'])
        for r in f['tablas']['Proveedores']:
            if str(r['codigo']).strip():propietarios[('proveedor',str(r['proveedor_id']),str(r['codigo']).strip().upper())].add(f['sku'])
    duplicados=set().union(*(v for v in propietarios.values() if len(v)>1)) if propietarios else set()
    for f in vista:
        if f['sku'] in duplicados:f['accion']='rechazado';f['errores'].append('La planilla asigna una misma barra o código del mismo proveedor a varios productos.')
    return vista


def firmar(vista,org,usuario,secreto):return URLSafeTimedSerializer(secreto,salt=FORMATO).dumps({'org':org,'usuario':usuario,'filas':vista})


def _guardar_atributos(filas,session,org):
    """Correcciones escalares ya validadas; no toca relaciones ni versiones de costo."""
    if not filas:return
    from sqlalchemy import update,bindparam
    from models.producto import Producto
    from models.producto_ficha_maestra import ProductoFichaMaestra
    from models.producto_datos_generales import ProductoDatosGenerales
    datos=[(f,leer_formulario_completo(f['adicionales'],[])) for f in filas]
    for Modelo,pk,campos,fuente in [(Producto,'id',('descripcion',),'base'),
            (ProductoFichaMaestra,'producto_id',('nombre','linea'),'base'),
            (ProductoDatosGenerales,'producto_id',('rubro','subrubro','marca','iva_pct','iibb_pct','utilidad_pct','utilidad_base'),'general')]:
        t=Modelo.__table__;valores={c:bindparam('valor_'+c) for c in campos}
        if Modelo is not Producto:valores['revision']=t.c.revision+1
        stmt=update(t).where(t.c[pk]==bindparam('_id'),t.c.organizacion_id==org).values(**valores)
        session.execute(stmt,[dict(_id=f['producto_id'],**{'valor_'+c:(f['datos'] if fuente=='base' else a)[c] for c in campos}) for f,a in datos])
    session.expire_all()


def aplicar(token,*,Producto,session,organizacion_id,usuario,secreto):
    from services.productos_ficha_maestra import guardar_ficha
    from services.productos_datos_generales import guardar_datos_generales
    from services.productos_produccion_maestro import preparar,calcular
    from services.productos_combo_maestro import preparar_composicion,guardar_composicion
    if len(token)>1000000:raise ValueError('La revisión supera el tamaño permitido.')
    try:lote=URLSafeTimedSerializer(secreto,salt=FORMATO).loads(token,max_age=900)
    except BadSignature:raise ValueError('La revisión venció o fue modificada; volvé a analizar la planilla.') from None
    if lote['org']!=organizacion_id or lote['usuario']!=usuario.id:raise ValueError('La revisión pertenece a otro usuario u organización.')
    filas=lote['filas']
    if not filas or any(f['errores'] for f in filas):raise ValueError('Corregí los errores antes de confirmar; no se guardó nada.')
    try:
        ids=[f['producto_id'] for f in filas]
        session.query(Producto).filter(Producto.organizacion_id==organizacion_id,Producto.id.in_(ids)).order_by(Producto.id).with_for_update().all()
        # Vuelve a validar identidad, huella, relaciones y costos en el momento de confirmar.
        tablas={h:[] for h in COLUMNAS}
        for f in filas:
            for h in COLUMNAS:tablas[h].extend(f['tablas'][h])
        revision=previsualizar(tablas,Producto=Producto,session=session,organizacion_id=organizacion_id)
        if any(f['errores'] for f in revision):raise ValueError('La planilla quedó desactualizada o dejó de ser válida. Volvé a analizarla.')
        originales={f['sku']:f for f in filas};cuenta={'actualizar':0,'sin_cambios':0}
        # Se validan TODAS las huellas y los importes antes de la primera escritura.
        for f in revision:
            if _firma(f['producciones'])!=_firma(originales[f['sku']]['producciones']):raise ValueError('Cambió el costo de fabricación desde la revisión; volvé a analizar.')
        _guardar_atributos([f for f in revision if f.get('solo_atributos')],session,organizacion_id)
        recetas_pendientes=[]
        for f in revision:
            if f.get('solo_atributos'):cuenta['actualizar']+=1;continue
            if f['accion']=='sin_cambios':cuenta['sin_cambios']+=1;continue
            solo_receta=bool(f['producciones']) and set(f['cambios'])<={'Perfiles','Componentes','Tareas','Consumos','Operaciones','Maquinas','Indirectos'}
            if solo_receta:
                p=session.get(Producto,f['producto_id'])
            else:
                p=guardar_ficha(f['datos'],Producto=Producto,db_session=session,organizacion_id=organizacion_id,producto_id=f['producto_id'],huella_esperada=f['huella'],commit=False)
                guardar_datos_generales(p,leer_formulario_completo(f['adicionales'],[]),session,usuario=usuario)
            for receta in f['producciones']:recetas_pendientes.append((p,receta))
            for combo in f['combos']:guardar_composicion(p,preparar_composicion(p,combo,session,organizacion_id),session)
            cuenta['actualizar']+=1
        from services.lectura_fabricacion_lote import LecturaFabricacion
        from services.guardar_recetas_planilla import guardar_recetas
        lectura=LecturaFabricacion(session,organizacion_id);preparadas=[]
        for p,r in recetas_pendientes:
            plan=preparar(r['receta'],session,organizacion_id,p.id,lectura=lectura);costo=calcular(plan,organizacion_id)
            if costo['huella']!=r['huella_costo']:raise ValueError('Cambió el costo al confirmar. Volvé a analizar.')
            preparadas.append((p,plan,costo))
        guardar_recetas(preparadas,session,organizacion_id,usuario)
        session.commit();return cuenta
    except Exception:session.rollback();raise
