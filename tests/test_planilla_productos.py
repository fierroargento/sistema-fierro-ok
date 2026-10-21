from io import BytesIO
from types import SimpleNamespace
from decimal import Decimal
import pytest
from openpyxl import load_workbook
from werkzeug.datastructures import FileStorage
from extensions import db
from models.producto import Producto
from models.producto_datos_generales import ProductoDatosGenerales
from models.costo_producto_version import CostoProductoVersion
from services import planilla_productos as servicio
from test_producto_ficha_maestra import entorno,crear
from test_producto_datos_generales import completo
from test_producto_fabricacion_maestro import fabricacion
from test_producto_despiece import despiece
from test_importacion_fabricacion_productos import preview,aplicar as importar_fabricacion


def archivo(d,editar=None):
    b=servicio.exportar(Producto=Producto,session=db.session,organizacion_id=d[0][1].id)
    if editar:
        w=load_workbook(b);editar(w);b=BytesIO();w.save(b);b.seek(0)
    return FileStorage(b,filename='productos.xlsx')


def revisar(d,a):return servicio.previsualizar(servicio.leer(a,d[0][1].id),Producto=Producto,session=db.session,organizacion_id=d[0][1].id)


def guardar(d,v):return servicio.aplicar(servicio.firmar(v,d[0][1].id,1,'clave'),Producto=Producto,session=db.session,organizacion_id=d[0][1].id,usuario=SimpleNamespace(id=1),secreto='clave')


def cambiar(w,hoja,columna,valor,fila=2):
    ws=w[hoja];cols=list(next(ws.values));ws.cell(fila,cols.index(servicio.etiqueta(columna))+1,valor)


def test_exportar_reimportar_sin_editar_cero_escrituras(despiece):
    from sqlalchemy import event
    importar_fabricacion(despiece,preview(despiece));antes=CostoProductoVersion.query.count();sql=[]
    def contar(conn,cursor,statement,*args):sql.append(statement.split()[0].upper())
    event.listen(db.engine,'before_cursor_execute',contar)
    try:
        v=revisar(despiece,archivo(despiece));assert [r['accion'] for r in v]==['sin_cambios'],v
        assert guardar(despiece,v)=={'actualizar':0,'sin_cambios':1}
        assert not set(sql)&{'INSERT','UPDATE','DELETE'}
    finally:event.remove(db.engine,'before_cursor_execute',contar)
    assert Producto.query.count()==1 and CostoProductoVersion.query.count()==antes


def test_editar_datos_iva_y_utilidad(despiece):
    importar_fabricacion(despiece,preview(despiece))
    def editar(w):
        cambiar(w,'Productos','nombre','Nombre corregido');cambiar(w,'Productos','iva_pct',10.5);cambiar(w,'Productos','utilidad_pct',120)
    v=revisar(despiece,archivo(despiece,editar));assert v[0]['errores']==[],v
    assert v[0]['accion']=='actualizar';guardar(despiece,v)
    from models.producto_ficha_maestra import ProductoFichaMaestra
    assert ProductoFichaMaestra.query.one().nombre=='Nombre corregido'
    a=ProductoDatosGenerales.query.one();assert a.iva_pct==Decimal('10.5') and a.utilidad_pct==120
    assert Producto.query.count()==CostoProductoVersion.query.count()==1
    assert revisar(despiece,archivo(despiece))[0]['accion']=='sin_cambios'


def test_editar_medida_recalcula_y_guarda_piezas(despiece):
    importar_fabricacion(despiece,preview(despiece))
    v=revisar(despiece,archivo(despiece,lambda w:cambiar(w,'Componentes','largo_cm',50)))
    assert not v[0]['errores'],v
    assert v[0]['producciones'][0]['costo_centavos']==10050
    guardar(despiece,v);assert Producto.query.count()==1 and CostoProductoVersion.query.count()==2
    from services.productos_ficha_maestra import construir_filas
    _,filas=construir_filas(Producto.query.all(),db.session,despiece[0][1].id)
    assert filas[0]['producciones'][0]['despiece']['actual']
    assert revisar(despiece,archivo(despiece))[0]['accion']=='sin_cambios'


def test_antigua_y_otra_org_rechazan(despiece):
    importar_fabricacion(despiece,preview(despiece));a=archivo(despiece)
    with pytest.raises(ValueError,match='organización'):servicio.leer(a,despiece[0][2].id)
    a=archivo(despiece);p=ProductoDatosGenerales.query.one();p.utilidad_pct=200;p.revision+=1;db.session.commit()
    assert revisar(despiece,a)[0]['errores']


def test_formula_y_hoja_faltante_rechazan(despiece):
    importar_fabricacion(despiece,preview(despiece))
    with pytest.raises(ValueError,match='fórmulas'):revisar(despiece,archivo(despiece,lambda w:cambiar(w,'Productos','nombre','=1+1')))
    with pytest.raises(ValueError,match='Falta la hoja'):revisar(despiece,archivo(despiece,lambda w:w.remove(w['Tareas'])))


def test_ruta_descarga_y_preview(despiece):
    importar_fabricacion(despiece,preview(despiece));cliente=despiece[0][0].test_client()
    r=cliente.get('/admin/productos/exportar.xlsx');assert r.status_code==200 and 'no-store' in r.headers['Cache-Control']
    r=cliente.post('/admin/productos',data={'accion':'previsualizar_planilla','_csrf_token':'csrf-test','archivo_planilla':(BytesIO(r.data),'productos.xlsx')})
    assert r.status_code==200 and '0 para actualizar' in r.text


def test_muchos_cambios_escalares_se_guardan_agrupados(despiece):
    from test_importacion_fabricacion_productos import lote
    from sqlalchemy import event
    importar_fabricacion(despiece,preview(despiece,lote(tuple(f'LOTE-{i}' for i in range(30)))))
    def editar(w):
        for i in range(2,32):cambiar(w,'Productos','utilidad_pct',150,i)
    v=revisar(despiece,archivo(despiece,editar));assert all(not f['errores'] and f['solo_atributos'] for f in v),v
    sql=[]
    def contar(conn,cursor,statement,*args):sql.append(statement.split()[0].upper())
    event.listen(db.engine,'before_cursor_execute',contar)
    try:guardar(despiece,v)
    finally:event.remove(db.engine,'before_cursor_execute',contar)
    assert sql.count('UPDATE')==3 and not set(sql)&{'INSERT','DELETE'}
    assert {r.utilidad_pct for r in ProductoDatosGenerales.query.all()}=={Decimal(150)}


def test_cambio_posterior_al_preview_y_rollback(despiece,monkeypatch):
    importar_fabricacion(despiece,preview(despiece))
    v=revisar(despiece,archivo(despiece,lambda w:cambiar(w,'Productos','nombre','Nuevo')))
    real=servicio._guardar_atributos
    def falla(*args):real(*args);raise ValueError('Fallo final')
    monkeypatch.setattr(servicio,'_guardar_atributos',falla)
    with pytest.raises(ValueError,match='Fallo final'):guardar(despiece,v)
    from models.producto_ficha_maestra import ProductoFichaMaestra
    assert ProductoFichaMaestra.query.one().nombre!='Nuevo'
    monkeypatch.undo()
    g=ProductoDatosGenerales.query.one();g.iva_pct=0;g.revision+=1;db.session.commit()
    with pytest.raises(ValueError,match='desactualizada'):guardar(despiece,v)


def test_simple_costo_moneda_y_barra_con_ceros(despiece):
    from test_producto_datos_generales import guardar as guardar_datos,formulario
    org=despiece[0][1];p=crear(org,sku='SIMPLE')
    guardar_datos(p,formulario(costo_maestro_presentado='1',precio_costo='100',moneda_costo='ARS',iva_pct='21',codigos_barras_json='["001234"]'))
    assert revisar(despiece,archivo(despiece))[0]['accion']=='sin_cambios'
    def editar(w):cambiar(w,'Productos','precio_costo',150);cambiar(w,'Productos','moneda_costo','USD')
    v=revisar(despiece,archivo(despiece,editar));assert not v[0]['errores'],v
    guardar(despiece,v)
    a=ProductoDatosGenerales.query.one();assert a.moneda_costo=='USD'
    from models.producto_datos_generales import ProductoCodigoBarras
    assert ProductoCodigoBarras.query.one().codigo=='001234'
    assert CostoProductoVersion.query.filter_by(moneda='USD').one().costo_total_centavos==15000


def test_combo_cantidad_propias_y_simple_preservado(despiece):
    org=despiece[0][1];crear(org,sku='BASE',habilita_componente=True)
    crear(org,sku='COMBO',tipo='combo',componentes_combo=[{'sku':'BASE','cantidad':'2'}])
    v=revisar(despiece,archivo(despiece,lambda w:cambiar(w,'Combos','cantidad',5)))
    assert not any(f['errores'] for f in v),v
    guardar(despiece,v)
    from models.perfil_costeo_producto import ComboProductoComponente
    assert ComboProductoComponente.query.one().cantidad==5 and Producto.query.count()==2
    assert all(f['accion']=='sin_cambios' for f in revisar(despiece,archivo(despiece)))


def test_exportacion_no_va_por_importador_basico(despiece):
    from services.importacion_maestro_productos import leer_lote
    importar_fabricacion(despiece,preview(despiece))
    with pytest.raises(ValueError,match='Actualizar desde planilla Fierro'):leer_lote(archivo(despiece))


def test_tarifa_cambiada_tras_preview_y_otra_firma(despiece):
    from models.fuentes_costo_productivo import InsumoPrecioVersion
    importar_fabricacion(despiece,preview(despiece))
    v=revisar(despiece,archivo(despiece,lambda w:cambiar(w,'Componentes','cantidad',3)))
    db.session.query(InsumoPrecioVersion).update({'precio_unitario_centavos':20000});db.session.commit()
    with pytest.raises(ValueError,match='costo'):guardar(despiece,v)
    assert CostoProductoVersion.query.count()==1
    kwargs=dict(Producto=Producto,session=db.session,organizacion_id=despiece[0][1].id,usuario=SimpleNamespace(id=1),secreto='clave')
    with pytest.raises(ValueError,match='usuario'):servicio.aplicar(servicio.firmar(v,despiece[0][1].id,2,'clave'),**kwargs)
    with pytest.raises(ValueError,match='modificada'):servicio.aplicar('inventado',**kwargs)


def test_omitir_producto_no_elimina_y_barra_duplicada_se_rechaza(despiece):
    org=despiece[0][1];crear(org,sku='A');crear(org,sku='B')
    def duplicar(w):w['Barras'].append(['A','00123']);w['Barras'].append(['B','00123'])
    v=revisar(despiece,archivo(despiece,duplicar));assert all(f['errores'] for f in v)
    with pytest.raises(ValueError):guardar(despiece,v)
    def omitir(w):w['Productos'].delete_rows(3)
    v=revisar(despiece,archivo(despiece,omitir));assert len(v)==1
    guardar(despiece,v);assert Producto.query.count()==2


def test_cambio_nombre_pieza_no_duplica_costo_y_fallo_final_revierte(despiece,monkeypatch):
    importar_fabricacion(despiece,preview(despiece))
    v=revisar(despiece,archivo(despiece,lambda w:cambiar(w,'Componentes','nombre','Pieza corregida')))
    guardar(despiece,v);assert CostoProductoVersion.query.count()==1
    from services.productos_ficha_maestra import construir_filas
    _,filas=construir_filas(Producto.query.all(),db.session,despiece[0][1].id)
    assert filas[0]['producciones'][0]['despiece']['datos']['piezas'][0]['nombre']=='Pieza corregida'
    v=revisar(despiece,archivo(despiece,lambda w:cambiar(w,'Componentes','largo_cm',70)))
    import services.guardar_recetas_planilla as persistencia
    real=persistencia._insertar
    def falla(session,tabla,filas,devolver=()):
        r=real(session,tabla,filas,devolver)
        if tabla.name=='costo_producto_detalle':raise ValueError('Fallo al final')
        return r
    monkeypatch.setattr(persistencia,'_insertar',falla)
    with pytest.raises(ValueError,match='Fallo al final'):guardar(despiece,v)
    assert CostoProductoVersion.query.count()==1
    _,filas=construir_filas(Producto.query.all(),db.session,despiece[0][1].id)
    assert Decimal(filas[0]['producciones'][0]['despiece']['datos']['piezas'][0]['largo_cm'])==40


def test_insumo_ajeno_rechazado_y_no_borra_perfil(despiece):
    importar_fabricacion(despiece,preview(despiece))
    v=revisar(despiece,archivo(despiece,lambda w:cambiar(w,'Componentes','insumo_id',despiece[0][-1].id)))
    assert v[0]['errores'] and Producto.query.count()==1
    v=revisar(despiece,archivo(despiece,lambda w:w['Perfiles'].delete_rows(2)))
    assert v[0]['errores']


def test_proveedores_imagenes_y_ubicaciones_se_conservan(despiece):
    from test_producto_datos_generales import guardar as guardar_datos,formulario,imagen
    from models.compras import ProveedorCompra
    from models.producto_datos_generales import ProductoProveedor,ProductoImagenMaestra
    import json
    p=crear(despiece[0][1],sku='CON-IMAGEN')
    proveedor=ProveedorCompra.query.filter_by(organizacion_id=p.organizacion_id).first()
    guardar_datos(p,formulario(proveedores_json=json.dumps([{'proveedor_id':proveedor.id,'codigo':'000123'}])),[imagen()])
    v=revisar(despiece,archivo(despiece));assert v[0]['accion']=='sin_cambios',v
    v=revisar(despiece,archivo(despiece,lambda w:cambiar(w,'Productos','nombre','Con imagen corregido')));guardar(despiece,v)
    assert ProductoProveedor.query.one().codigo=='000123' and ProductoImagenMaestra.query.count()==1


def test_muchas_recetas_editadas_en_grupo(despiece):
    from test_importacion_fabricacion_productos import lote
    from sqlalchemy import event
    importar_fabricacion(despiece,preview(despiece,lote(tuple(f'RECETA-{i}' for i in range(30)))))
    def editar(w):
        for i in range(2,32):cambiar(w,'Componentes','largo_cm',55,i)
    v=revisar(despiece,archivo(despiece,editar));assert all(not f['errores'] for f in v),v
    sql=[]
    def contar(conn,cursor,statement,*args):sql.append(statement.split()[0].upper())
    event.listen(db.engine,'before_cursor_execute',contar)
    try:guardar(despiece,v)
    finally:event.remove(db.engine,'before_cursor_execute',contar)
    assert sql.count('UPDATE')<10 and sql.count('INSERT')<10 and sql.count('DELETE')<=4
    assert Producto.query.count()==30 and CostoProductoVersion.query.count()==60


def test_ficha_anterior_con_consumos_operaciones_maquina_indirecto(despiece):
    from test_producto_fabricacion_maestro import alta
    app=despiece[0][0];assert app.test_client().post('/admin/productos',data=alta(despiece[0])).status_code==302
    assert revisar(despiece,archivo(despiece))[0]['accion']=='sin_cambios'
    v=revisar(despiece,archivo(despiece,lambda w:cambiar(w,'Consumos','cantidad',3)))
    assert not v[0]['errores'],v
    guardar(despiece,v)
    assert CostoProductoVersion.query.count()==2
    assert revisar(despiece,archivo(despiece))[0]['accion']=='sin_cambios'


def test_dos_unidades_conservan_identidad_y_solo_edita_una(despiece):
    from test_importacion_fabricacion_productos import lote,opciones
    importar_fabricacion(despiece,preview(despiece))
    from services.productos_produccion_maestro import preparar,guardar as guardar_produccion
    p=Producto.query.one();receta=preview(despiece,lote(('SEGUNDO',)))['filas'][0]['produccion'];receta['unidad_id']=despiece[0][3].id
    plan=preparar(receta,db.session,p.organizacion_id,p.id);guardar_produccion(p,plan,db.session);db.session.commit()
    v=revisar(despiece,archivo(despiece));assert v[0]['accion']=='sin_cambios',v
    def editar(w):cambiar(w,'Componentes','cantidad',4,3)
    v=revisar(despiece,archivo(despiece,editar));assert not v[0]['errores'],v
    guardar(despiece,v)
    from models.perfil_costeo_producto import PerfilCosteoProducto
    assert PerfilCosteoProducto.query.count()==2 and Producto.query.count()==1
    assert revisar(despiece,archivo(despiece))[0]['accion']=='sin_cambios'


def test_texto_exportado_no_es_formula_y_no_filtra_otra_org(despiece):
    crear(despiece[0][1],sku='SEGURIDAD',nombre='=texto literal')
    crear(despiece[0][2],sku='AJENO-PRIVADO',nombre='Dato de otra empresa')
    a=archivo(despiece);wb=load_workbook(a.stream)
    nombres=list(next(wb['Productos'].values));celda=wb['Productos'].cell(2,nombres.index('Nombre')+1)
    assert celda.value=='=texto literal' and celda.data_type=='s'
    assert all('AJENO-PRIVADO' not in str(v) for ws in wb for row in ws.values for v in row)
    a.stream.seek(0);v=revisar(despiece,a);assert v[0]['accion']=='sin_cambios',v[0]['detalle_cambios']


def test_completar_tipo_unidad_de_producto_anterior_sin_duplicar(completo):
    from models.producto_ficha_maestra import ProductoFichaMaestra
    org=completo[1]
    p=Producto(organizacion_id=org.id,sku='UAT-MP-001',descripcion='Materia prima anterior')
    db.session.add(p);db.session.commit();ident=p.id
    d=(completo,)
    a=archivo(d)
    w=load_workbook(a.stream)
    ws=w['Productos'];cols=list(next(ws.values))
    assert ws.cell(2,cols.index('Tipo')+1).value is None
    assert not ws.cell(2,cols.index('Tipo')+1).protection.locked
    a.stream.seek(0)
    assert revisar(d,a)[0]['accion']=='sin_cambios'
    def editar(w):
        cambiar(w,'Productos','tipo','Simple');cambiar(w,'Productos','unidad_medida','unidad');cambiar(w,'Productos','marca','Fierro')
    v=revisar(d,archivo(d,editar))
    assert not v[0]['errores'],v
    assert v[0]['accion']=='actualizar' and not v[0]['solo_atributos']
    assert ProductoFichaMaestra.query.count()==0
    assert guardar(d,v)=={'actualizar':1,'sin_cambios':0}
    assert Producto.query.one().id==ident
    f=ProductoFichaMaestra.query.one()
    assert f.tipo=='simple' and f.unidad_medida=='unidad'
    assert ProductoDatosGenerales.query.one().marca=='Fierro'
    assert revisar(d,archivo(d))[0]['accion']=='sin_cambios'


def test_ficha_incompleta_indica_columnas_y_resumen_visible(completo):
    org=completo[1]
    db.session.add(Producto(organizacion_id=org.id,sku='UAT-MP-001',descripcion='Anterior'));db.session.commit()
    d=(completo,)
    a=archivo(d,lambda w:cambiar(w,'Productos','marca','Fierro'))
    v=revisar(d,a)
    assert '«Tipo»' in v[0]['errores'][0] and '«Unidad de medida»' in v[0]['errores'][0]
    a=archivo(d,lambda w:cambiar(w,'Productos','marca','Fierro'))
    from test_producto_ficha_maestra import configurar_rutas
    configurar_rutas(completo[0],org)
    r=completo[0].test_client().post('/admin/productos',data={'accion':'previsualizar_planilla','_csrf_token':'csrf-test','archivo_planilla':(a.stream,'fierro.xlsx')})
    assert r.status_code==200
    import re
    aviso=re.search(r'<div class="productos-preview" role="alert">(.*?)</div>',r.text,re.S)
    assert aviso and 'UAT-MP-001' in aviso[1] and 'fila 2' in aviso[1]
    confirmacion=r.text.split('name="planilla_token"',1)[1].split('</form>',1)[0]
    assert re.search(r'<button[^>]*disabled',confirmacion)


def test_tipo_y_vinculo_existentes_siguen_protegidos(completo):
    d=(completo,);crear(completo[1])
    for campo,valor,texto in [('tipo','combo','Columna «Tipo»'),('insumo_id',123,'Columna «Insumo ID»')]:
        v=revisar(d,archivo(d,lambda w:cambiar(w,'Productos',campo,valor)))
        assert texto in v[0]['errores'][0]
        with pytest.raises(ValueError,match='Corregí'):guardar(d,v)
    assert Producto.query.count()==1


@pytest.mark.parametrize('tipo',['combo','produccion'])
def test_inicializar_composicion_sin_componentes_no_se_permite(completo,tipo):
    db.session.add(Producto(organizacion_id=completo[1].id,sku='ANTERIOR',descripcion='Anterior'));db.session.commit()
    def editar(w):
        cambiar(w,'Productos','tipo',tipo);cambiar(w,'Productos','unidad_medida','unidad')
    v=revisar((completo,),archivo((completo,),editar))
    assert v[0]['errores'] and v[0]['accion']=='rechazado'


def test_resumen_identifica_fila_real_si_hay_blancos(completo):
    db.session.add(Producto(organizacion_id=completo[1].id,sku='ANTERIOR',descripcion='Anterior'));db.session.commit()
    def editar(w):
        w['Productos'].insert_rows(2,3)
        cambiar(w,'Productos','marca','Fierro',fila=5)
    v=revisar((completo,),archivo((completo,),editar))
    assert v[0]['numero']==5 and v[0]['errores']
