"""Lista de proveedor: revisión sin escrituras, unidad, historia y costo explícito."""
from io import BytesIO
from types import SimpleNamespace
from decimal import Decimal
import pytest
from werkzeug.datastructures import FileStorage,MultiDict
from extensions import db
from models.organizacion import Organizacion
from models.unidad_negocio import UnidadNegocio
from models.existencia_sucursal import ExistenciaSucursal
from models.inventario_saas import ItemInventario
from models.sucursal_operativa import SucursalOperativa
from models.compras import ProveedorCompra as P
from models.proveedor_unidad import ProveedorUnidad as PU
from models.fuentes_costo_productivo import InsumoProductivo as I,InsumoPrecioVersion as V
from models.listas_proveedores_compra import MaterialProveedorCompra as L,CotizacionProveedorCompra as C
from services import listas_proveedores_compra as s
from scripts.preparar_listas_proveedores import preparar,TABLAS
from test_personal_admin import personal,entorno

@pytest.fixture()
def listas(personal,monkeypatch):
    import modules.admin.compras.routes as routes
    c,org,otra,a,m,mem=personal
    monkeypatch.setattr(routes,'resolver_tenant_usuario',lambda *args,**kw:mem)
    db.metadata.create_all(db.engine,tables=[P.__table__,PU.__table__,V.__table__])
    b=UnidadNegocio(organizacion_id=org.id,codigo='nautica-listas',nombre='Náutica',activa=True)
    ps=[P(organizacion_id=o,codigo=cod,razon_social=nom,estado='activo') for o,cod,nom in [(org.id,'PA','Proveedor A'),(org.id,'PB','Proveedor B'),(otra.id,'PX','Proveedor ajeno'),(org.id,'PO','Oculto en otra unidad')]]
    db.session.add_all([b,*ps]);db.session.flush()
    for p,u in [(ps[0],a),(ps[0],b),(ps[1],a),(ps[3],b)]:db.session.add(PU(organizacion_id=org.id,unidad_negocio_id=u.id,proveedor_id=p.id,estado='activo'))
    items=[I(organizacion_id=o,unidad_negocio_id=u,codigo=cod,nombre=nom,tipo='materia_prima',unidad_medida='metro',activo=True) for o,u,cod,nom in [(org.id,a.id,'MAT-A','Material A'),(org.id,a.id,'MAT-A2','Material A2'),(org.id,b.id,'MAT-B','Material Náutica'),(org.id,None,'GENERAL','Material general'),(otra.id,None,'AJENO','Material ajeno')]]
    db.session.add_all(items);db.session.commit();preparar(db.engine);preparar(db.engine)
    m.update(Organizacion=Organizacion,InsumoProductivo=I,InsumoPrecioVersion=V,ProveedorCompra=P)
    c.application.register_blueprint(routes.crear_blueprint_compras(dependencias=dict(db=db,modelos=m,usuario_actual=lambda:SimpleNamespace(id=1),UsuarioOrganizacion=m['UsuarioOrganizacion'],login_required=lambda f:f,registrar_auditoria=lambda *args,**kw:None)))
    from services.normalizacion_datos_formulario import registrar
    registrar(c.application)
    yield c,org,otra,a,b,ps,items,mem
    from models.insumo_presentacion_compra import InsumoPrecioBultoVersion as BP,InsumoPresentacionCompra as IP
    db.session.rollback();db.metadata.drop_all(db.engine,tables=[BP.__table__,IP.__table__]+list(reversed(TABLAS))+[V.__table__,PU.__table__,P.__table__])

def fila(codigo='EXT-001',precio='600',base='bulto',cantidad='6',moneda='ARS',**kw):
    return dict(numero=2,valores=dict(codigo_proveedor=codigo,precio_proveedor=precio,base_precio=base,cantidad_bulto=cantidad,moneda_proveedor=moneda,medidas_origen='Barra de 6 m')|kw,formulas=[])

def vinculo(datos,indice=0,codigo='EXT-001',proveedor=0,unidad=None):
    _,org,_,a,b,ps,items,*_=datos
    l=s.vincular(items[indice],ps[proveedor],codigo,org.id,(unidad or a).id,db.session);db.session.commit();return l

def revision(datos,filas=None,campos=None,proveedor=0,unidad=None):
    _,org,_,a,b,ps,*_=datos
    return s.revisar(filas or [fila()],campos or list(s.CAMPOS),ps[proveedor],org.id,(unidad or a).id,db.session)

def aplicar(datos,vista,proveedor=0,unidad=None,usuario=1):
    _,org,_,a,b,ps,*_=datos;u=unidad or a
    token=s.firmar(vista,ps[proveedor].id,org.id,u.id,usuario,'secreto')
    n=s.confirmar(token,ps[proveedor],org.id,u.id,usuario,'secreto',db.session);db.session.commit();return n,token

def form(datos,accion,**kw):
    _,_,_,a,_,ps,*_=datos
    return dict(_csrf_token='csrf-personal',unidad_contexto_id=str(a.id),proveedor_id=str(ps[0].id),accion=accion)|kw

def test_importa_con_historia_y_costo_solo_al_elegirlo(listas):
    l=vinculo(listas);v=revision(listas)
    assert v['filas'][0]['plan']['precio_unitario_centavos']==10000
    assert C.query.count()==0 and V.query.count()==0
    assert aplicar(listas,v)[0]==1
    c=C.query.one();assert c.precio_bulto_centavos==60000 and c.cantidad_bulto==6 and V.query.count()==0
    # Guardar una cotización de bulto no exige aún ni modifica la presentación del material.
    from models.insumo_presentacion_compra import InsumoPresentacionCompra as IP,InsumoPrecioBultoVersion as BP
    db.metadata.create_all(db.engine,tables=[IP.__table__,BP.__table__])
    try:
        _,org,_,a,_,ps,*_=listas
        fuente=s.usar_en_costo(c.id,ps[0],org.id,a.id,'',db.session,None);db.session.commit()
        assert fuente.precio_unitario_centavos==10000 and fuente.proveedor_referencia==ps[0].razon_social
        assert db.session.get(BP,fuente.id).precio_bulto_centavos==60000
        assert V.query.count()==1
        s.usar_en_costo(c.id,ps[0],org.id,a.id,fuente.id,db.session,None);db.session.commit()
        assert V.query.count()==1
        vista=revision(listas,[fila(precio='720')]);assert aplicar(listas,vista)[0]==1
        db.session.expire_all();assert C.query.count()==2 and not db.session.get(C,c.id).vigente
        assert V.query.filter_by(vigente=True).one().precio_unitario_centavos==10000
    finally:
        db.session.rollback();db.metadata.drop_all(db.engine,tables=[BP.__table__,IP.__table__])

def test_proveedores_unidades_y_monedas_independientes(listas):
    _,org,_,a,b,ps,items,*_=listas
    vinculo(listas);vinculo(listas,proveedor=1);vinculo(listas,indice=2,unidad=b)
    aplicar(listas,revision(listas,[fila(base='unitario',precio='10',cantidad='')]))
    aplicar(listas,revision(listas,[fila(base='unitario',precio='20',cantidad='')],proveedor=1),proveedor=1)
    aplicar(listas,revision(listas,[fila(base='unitario',precio='30',cantidad='')],unidad=b),unidad=b)
    aplicar(listas,revision(listas,[fila(base='unitario',precio='1',cantidad='',moneda='USD')]))
    assert sorted(c.precio_unitario_centavos for c in C.query.filter_by(vigente=True))==[100,1000,2000,3000]
    assert V.query.count()==0
    with pytest.raises(ValueError):s.vincular(items[2],ps[0],'OTRO',org.id,a.id,db.session)
    with pytest.raises(ValueError):s.vincular(items[3],ps[0],'GENERAL',org.id,a.id,db.session)
    with pytest.raises(ValueError):s.vincular(items[0],ps[2],'AJENO',org.id,a.id,db.session)

def test_ignora_no_vinculados_no_crea_productos_ni_aplica_formulas_ajenas(listas):
    vinculo(listas)
    x=fila('NO-COMPRAMOS');x['formulas']=['precio_proveedor']
    v=revision(listas,[fila(),x,fila(codigo='OTRO',sku='MAT-A')])
    assert v['ignoradas']==2 and len(v['filas'])==1
    n0=I.query.count();assert aplicar(listas,v)[0]==1 and I.query.count()==n0
    assert L.query.count()==1

def test_campos_vacios_conservan_y_cero_se_aplica(listas):
    vinculo(listas);aplicar(listas,revision(listas))
    v=revision(listas,[fila(precio='',base='',cantidad='99')],campos=['precio_proveedor'])
    assert not v['filas'][0]['cambia'] and not v['filas'][0]['errores']
    v=revision(listas,[fila(precio='0',cantidad='999')],campos=['precio_proveedor'])
    assert v['filas'][0]['plan']['cantidad_bulto']=='6.000000'
    aplicar(listas,v);assert C.query.filter_by(vigente=True).one().precio_unitario_centavos==0
    assert C.query.filter_by(vigente=True).one().cantidad_bulto==6

def test_datos_invalidos_y_duplicados_no_se_confirman(listas):
    vinculo(listas)
    for filas in [[fila(),fila()],[fila(cantidad='0')],[fila(precio='-1')],[fila(sku='OTRO')],[dict(fila(),formulas=['precio_proveedor'])]]:
        v=revision(listas,filas)
        assert any(f['errores'] for f in v['filas'])
        with pytest.raises(ValueError):s.firmar(v,listas[5][0].id,listas[1].id,listas[3].id,1,'secreto')
    assert C.query.count()==0

def test_revision_firmada_detecta_cambios_usuario_y_contexto(listas):
    _,org,_,a,b,ps,items,*_=listas
    l=vinculo(listas);v=revision(listas);token=s.firmar(v,ps[0].id,org.id,a.id,1,'secreto')
    for p,u,usuario in [(ps[1],a,1),(ps[0],b,1),(ps[0],a,2)]:
        with pytest.raises(ValueError):s.confirmar(token,p,org.id,u.id,usuario,'secreto',db.session)
    with pytest.raises(ValueError):s.confirmar(token+'x',ps[0],org.id,a.id,1,'secreto',db.session)
    items[0].nombre='Cambió';db.session.commit()
    with pytest.raises(ValueError,match='Cambió'):s.confirmar(token,ps[0],org.id,a.id,1,'secreto',db.session)
    assert C.query.count()==0
    v=revision(listas);_,token=aplicar(listas,v)
    with pytest.raises(ValueError):s.confirmar(token,ps[0],org.id,a.id,1,'secreto',db.session)
    assert C.query.count()==1

def test_xlsx_csv_reutiliza_lector_y_plantilla(listas):
    vinculo(listas)
    csv='Código proveedor;Precio proveedor sin IVA;Cotización proveedor;Cantidad por bulto;Moneda proveedor;Medidas en origen\nEXT-001;600;bulto;6;ARS;barra\nDESCARTAR;999;unitario;;ARS;\n'
    filas,campos=s.lector.leer(FileStorage(stream=BytesIO(csv.encode()),filename='lista.csv'),list(s.CAMPOS))
    v=revision(listas,filas,campos);assert v['ignoradas']==1 and aplicar(listas,v)[0]==1
    from openpyxl import load_workbook
    w=load_workbook(s.plantilla());assert w['Productos'].cell(1,1).value=='Código proveedor'
    assert all(c.value is None for row in w['Productos'].iter_rows(min_row=2) for c in row)
    assert w['Productos']['A2'].number_format=='@';w.close()

def test_rutas_carga_revision_confirmacion_y_protecciones(listas):
    c,org,_,a,b,ps,items,mem=listas
    r=c.get('/admin/compras/listas-proveedores');assert r.status_code==200,r.text
    assert 'Listas de proveedores' in r.text and items[2].nombre not in r.text
    assert c.get('/admin/compras/listas-proveedores?proveedor_id='+str(ps[3].id)).status_code==404
    assert c.get('/admin/compras/listas-proveedores/plantilla').status_code==200
    r=c.post('/admin/compras/listas-proveedores',data=form(listas,'vincular_material',insumo_id=items[0].id,codigo_proveedor='EXT-001'))
    assert r.status_code==302,r.text
    r=c.post('/admin/compras/listas-proveedores',data=form(listas,'revisar_manual',vinculo_id=L.query.one().id,precio_proveedor='12',base_precio='unitario',moneda_proveedor='ARS'))
    assert r.status_code==200 and 'Revisar cambios' in r.text,r.text
    assert C.query.count()==0
    from html.parser import HTMLParser
    class Token(HTMLParser):
        token=''
        def handle_starttag(self,tag,attrs):
            a=dict(attrs)
            if tag=='input' and a.get('name')=='token_lista':self.token=a['value']
    t=Token();t.feed(r.text);assert t.token
    r=c.post('/admin/compras/listas-proveedores',data=form(listas,'confirmar_lista',token_lista=t.token))
    assert r.status_code==302 and C.query.count()==1
    r=c.get(r.location);assert r.status_code==200 and 'Usar en costo' in r.text
    assert c.get('/admin/compras/listas-proveedores?proveedor_id='+str(ps[1].id)+'&ver='+str(L.query.one().id)).status_code==404
    assert c.post('/admin/compras/listas-proveedores',data=form(listas,'vincular_material',insumo_id=items[1].id,codigo_proveedor='X',unidad_contexto_id=b.id)).status_code==400
    f=form(listas,'vincular_material');f.pop('_csrf_token');assert c.post('/admin/compras/listas-proveedores',data=f).status_code==400
    mem.rol='carga';assert c.get('/admin/compras/listas-proveedores').status_code==302

def test_sin_preparacion_conserva_datos_y_no_escribe(listas):
    c,org,_,a,b,ps,items,*_=listas
    db.metadata.drop_all(db.engine,tables=list(reversed(TABLAS)))
    try:
        r=c.get('/admin/compras/listas-proveedores?proveedor_id='+str(ps[0].id));assert r.status_code==200 and 'preparación de listas' in r.text
        assert c.post('/admin/compras/listas-proveedores',data=form(listas,'vincular_material',insumo_id=items[0].id,codigo_proveedor='X')).status_code==400
        assert I.query.count()==5
    finally:preparar(db.engine)


def test_reutiliza_vinculo_de_producto_con_fuente_de_unidad(listas):
    from models.producto import Producto
    from models.producto_ficha_maestra import ProductoFichaMaestra as F
    from models.producto_datos_generales import ProductoProveedor as PP
    c,org,_,a,b,ps,items,*_=listas
    db.metadata.create_all(db.engine,tables=[F.__table__,PP.__table__])
    try:
        productos=[Producto(organizacion_id=org.id,sku=sku,descripcion='Producto de prueba') for sku in ['SKU-VIEJO','SKU-OTRA-UNIDAD','SIN-FUENTE']]
        db.session.add_all(productos);db.session.flush()
        for p,i in zip(productos,[items[0],items[2],None]):
            db.session.add(F(producto_id=p.id,organizacion_id=org.id,nombre=p.sku,linea='',unidad_medida='metro',tipo='simple',activo=True,habilita_compra=True,habilita_venta=True,habilita_insumo=True,habilita_componente=False,insumo_id=i.id if i else None))
            db.session.add(PP(organizacion_id=org.id,producto_id=p.id,proveedor_id=ps[0].id,codigo=p.sku,codigo_clave=p.sku,principal=True))
        db.session.commit()
        vista=revision(listas,[fila(codigo='SKU-VIEJO',sku='SKU-VIEJO'),fila(codigo='SKU-OTRA-UNIDAD'),fila(codigo='SIN-FUENTE')])
        assert vista['ignoradas']==2 and len(vista['filas'])==1
        assert L.query.count()==0  # Previsualizar no convierte el vínculo ni escribe.
        assert aplicar(listas,vista)[0]==1
        l=L.query.one();assert l.insumo_id==items[0].id and l.codigo=='SKU-VIEJO'
        assert Producto.query.count()==3 and I.query.count()==5
        l.codigo='NUEVO-CODIGO';l.codigo_clave='NUEVO-CODIGO';db.session.commit()
        assert revision(listas,[fila(codigo='SKU-VIEJO')])['ignoradas']==1
    finally:
        db.session.rollback();db.metadata.drop_all(db.engine,tables=[PP.__table__,F.__table__])


def test_costo_ajeno_historico_o_desactualizado_no_aplica(listas):
    _,org,_,a,b,ps,items,*_=listas
    vinculo(listas);aplicar(listas,revision(listas));anterior=C.query.one()
    aplicar(listas,revision(listas,[fila(precio='900')]))
    for cid,p,u in [(anterior.id,ps[0],a),(C.query.filter_by(vigente=True).one().id,ps[1],a),(C.query.filter_by(vigente=True).one().id,ps[0],b)]:
        with pytest.raises(ValueError):s.usar_en_costo(cid,p,org.id,u.id,'',db.session,None)
    assert V.query.count()==0
    vigente=C.query.filter_by(vigente=True).one()
    with pytest.raises(ValueError):s.usar_en_costo(vigente.id,ps[0],org.id,a.id,'999',db.session,None)
    assert V.query.count()==0


def test_pantalla_historia_manual_e_importacion_bloquean_errores_y_contexto(listas):
    from html.parser import HTMLParser
    c,org,_,a,b,ps,items,mem=listas
    vinculo(listas);aplicar(listas,revision(listas));l=L.query.one()
    r=c.get(f'/admin/compras/listas-proveedores?proveedor_id={ps[0].id}&ver={l.id}')
    assert r.status_code==200 and 'lista-historia' in r.text
    assert 'lista-gestionar-'+str(l.id) in r.text
    class Forms(HTMLParser):
        def __init__(self):super().__init__();self.dentro=False;self.anidados=0
        def handle_starttag(self,tag,attrs):
            if tag=='form':self.anidados+=int(self.dentro);self.dentro=True
        def handle_endtag(self,tag):
            if tag=='form':self.dentro=False
    html=Forms();html.feed(r.text);assert html.anidados==0
    assert C.query.count()==1  # Abrir historial no modifica.
    csv='Código proveedor;Precio proveedor sin IVA;Cotización proveedor;Moneda proveedor\nEXT-001;180;unitario;ARS\n'
    data=MultiDict(form(listas,'revisar_lista'));data.setlist('campos_actualizar',['precio_proveedor','base_precio','moneda_proveedor'])
    data['archivo_lista']=(BytesIO(csv.encode()),'lista.csv')
    r=c.post('/admin/compras/listas-proveedores',data=data,content_type='multipart/form-data')
    assert r.status_code==200 and 'Confirmar cotizaciones' in r.text,r.text
    assert C.query.count()==1 and V.query.count()==0
    r=c.post('/admin/compras/listas-proveedores',data=form(listas,'revisar_manual',vinculo_id=l.id,precio_proveedor='-1',base_precio='unitario',moneda_proveedor='ARS'))
    assert r.status_code==200 and 'token_lista" value=' not in r.text
    assert C.query.count()==1
    # Código de proveedor editable sin reemplazar el material ni sus versiones.
    r=c.post('/admin/compras/listas-proveedores',data=form(listas,'editar_codigo',vinculo_id=l.id,codigo_anterior=l.codigo,codigo_proveedor='NUEVO'))
    assert r.status_code==302 and L.query.one().insumo_id==items[0].id and C.query.count()==1
    assert c.post('/admin/compras/listas-proveedores',data=form(listas,'editar_codigo',vinculo_id=l.id,codigo_anterior='EXT-001',codigo_proveedor='OBSOLETO')).status_code==400
    # Una lista válida pero de otra unidad no se confirma al cambiar la selección.
    with c.session_transaction() as ses:ses['unidad_negocio_id']=b.id
    r=c.post('/admin/compras/listas-proveedores',data=form(listas,'usar_costo',cotizacion_id=C.query.one().id,version_fuente=''))
    assert r.status_code==400 and V.query.count()==0


def test_proveedor_inactivo_solo_consulta_y_no_actualiza(listas):
    c,org,_,a,b,ps,items,*_=listas
    vinculo(listas);aplicar(listas,revision(listas))
    pu=PU.query.filter_by(proveedor_id=ps[0].id,unidad_negocio_id=a.id).one();pu.estado='inactivo';db.session.commit()
    r=c.get('/admin/compras/listas-proveedores?proveedor_id='+str(ps[0].id))
    assert r.status_code==200 and 'proveedor está inactivo' in r.text
    with pytest.raises(ValueError):revision(listas)
    r=c.post('/admin/compras/listas-proveedores',data=form(listas,'usar_costo',cotizacion_id=C.query.one().id,version_fuente=''))
    assert r.status_code==400 and V.query.count()==0 and C.query.count()==1


def test_preparacion_cli_aislada_sin_iniciar_app(tmp_path):
    import os,sys,sqlite3,subprocess
    archivo=tmp_path/'preparacion.db'
    with sqlite3.connect(archivo) as conn:
        for tabla in ['organizacion','unidad_negocio','proveedor_compra','insumo_productivo','insumo_precio_version']:
            conn.execute('CREATE TABLE '+tabla+' (id INTEGER PRIMARY KEY)')
    env=os.environ.copy();env.update(DATABASE_URL='sqlite:///'+str(archivo),MODO_LABORATORIO_DESCONECTADO='true',RENDER_GIT_BRANCH='integracion-saas-2026-09',RENDER_EXTERNAL_URL='https://sistema-fierro-crm-staging.onrender.com')
    for _ in range(2):
        r=subprocess.run([sys.executable,'-m','scripts.preparar_listas_proveedores','--staging'],env=env,capture_output=True,text=True)
        assert r.returncode==0,r.stderr
        assert 'Sin modificar materiales' in r.stdout
    with sqlite3.connect(archivo) as conn:
        nombres={r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        assert {'material_proveedor_compra','cotizacion_proveedor_compra','insumo_presentacion_compra','insumo_precio_bulto_version'} <= nombres
    env['MODO_LABORATORIO_DESCONECTADO']='false'
    r=subprocess.run([sys.executable,'-m','scripts.preparar_listas_proveedores','--staging'],env=env,capture_output=True,text=True)
    assert r.returncode!=0 and 'laboratorio desconectado' in r.stderr


def test_importes_locales_sin_ambiguedad_ni_valores_no_finitos(listas):
    vinculo(listas)
    v=revision(listas,[fila(precio='1.234,56',base='unitario',cantidad='')])
    assert v['filas'][0]['plan']['precio_unitario_centavos']==123456
    for precio in ['1.500','NaN','Infinity','texto']:
        v=revision(listas,[fila(precio=precio,base='unitario',cantidad='')])
        assert v['filas'][0]['errores']
    assert C.query.count()==0
