from datetime import datetime
from io import BytesIO
from zipfile import ZipFile
from openpyxl import load_workbook
from extensions import db
from models.personal_administrativo import PersonalDetalle,PersonalAcceso
from models.personas import CompetenciaEmpleadoPreparatoria,TurnoLaboralPlanificado
from services.personal_legajo import guardar_horario,guardar_ausencia
from services.cierre_personas_preparatorio import registrar_competencia
from test_personal_admin import personal,entorno,nuevo
from test_personal_actualizacion_costos import sueldo
from test_personal_semana_editor import fila


def libro(c,**kw):
    r=c.get('/admin/personal/exportar',query_string=kw)
    assert r.status_code==200 and r.mimetype=='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'attachment;' in r.headers['Content-Disposition'] and '.xlsx' in r.headers['Content-Disposition']
    assert 'no-store' in r.headers['Cache-Control']
    return load_workbook(BytesIO(r.data)),r.data


def test_exporta_legajo_y_todas_las_secciones_con_historial_y_sin_secretos(personal):
    c,org,_,u,m,*_=personal;e=nuevo(personal)
    sueldo(e);sueldo(e)
    db.session.add(PersonalDetalle(empleado_id=e.id,organizacion_id=org.id,dni='00000001',telefono='0012345',email='persona@example.test'))
    h=guardar_horario(e,fila(),db.session);h.estado='cancelado';db.session.flush()
    guardar_horario(e,fila(0,'16:00','20:00'),db.session)
    guardar_ausencia(e,dict(tipo='vacaciones',desde='2027-01-04',hasta='2027-01-08',detalle='Vacaciones de prueba',periodo='2026',dias_computados='5'),db.session)
    db.session.add(TurnoLaboralPlanificado(empleado_id=e.id,organizacion_id=org.id,unidad_negocio_id=u.id,
        inicio=datetime(2026,10,26,8),fin=datetime(2026,10,26,16),sector='Taller',clave_idempotencia='test-export'))
    user=m['UsuarioSistema'](username='usuario-export',nombre='Persona',password_hash='NO_EXPORTAR_HASH',rol='carga',activo=True)
    db.session.add(user);db.session.flush()
    mem=m['UsuarioOrganizacion'](usuario_id=user.id,organizacion_id=org.id,rol='carga',activa=True)
    db.session.add(mem);db.session.flush()
    db.session.add(PersonalAcceso(empleado_id=e.id,usuario_organizacion_id=mem.id,codigo_alarma_cifrado=b'NO_EXPORTAR_CIFRADO'))
    registrar_competencia(dict(codigo='SOLD',nombre='Soldadura',nivel='3'),empleado=e,organizacion_id=org.id,
        unidad_negocio_id=u.id,Competencia=CompetenciaEmpleadoPreparatoria,db_session=db.session)
    antes=db.session.query(m['EmpleadoCostoVersion']).count()
    wb,datos=libro(c)
    assert set(wb.sheetnames)=={'Personal','Sueldos','Horarios','Vacaciones y licencias','Novedades','Turnos','Competencias','Distribución de costos','Equipos','Áreas mano de obra','Accesos','Información'}
    personal=list(wb['Personal'].values);assert personal[1][6]=='00000001' and personal[1][11]=='0012345'
    assert wb['Sueldos'].max_row==3 and wb['Horarios'].max_row==3
    assert {r[9] for r in list(wb['Horarios'].values)[1:]}=={'cancelado','vigente'}
    assert wb['Vacaciones y licencias'].max_row==2 and wb['Novedades'].max_row==2
    assert wb['Turnos'].max_row==2 and wb['Competencias'].max_row==2
    assert list(wb['Accesos'].values)[1][2:] == ('usuario-export','carga',True,True)
    assert wb['Sueldos']['H2'].value==1000 and wb['Sueldos']['H2'].number_format=='#,##0.00'
    with ZipFile(BytesIO(datos)) as z:
        xml=b''.join(z.read(p) for p in z.namelist() if p.endswith('.xml'))
        assert b'NO_EXPORTAR_HASH' not in xml and b'NO_EXPORTAR_CIFRADO' not in xml
    assert db.session.query(m['EmpleadoCostoVersion']).count()==antes
    assert wb['Personal'].freeze_panes=='A2' and wb['Personal'].auto_filter.ref
    wb.close()


def test_exportacion_no_depende_de_filtros_pagina_o_estado_y_excluye_recursos_y_otros_tenants(personal):
    c,org,otra,u,m,*_=personal
    for i in range(27):
        e=nuevo(personal,codigo=f'EMP-{i:03d}')
        if i==26:e.activo=False;db.session.commit()
    recurso=nuevo(personal,codigo='equipo');recurso.tipo_registro='recurso'
    ajeno=nuevo(personal,codigo='ajeno');ajeno.organizacion_id=otra.id
    db.session.commit()
    wb,_=libro(c,buscar='no-coincide',pagina='2')
    filas=list(wb['Personal'].values)[1:]
    assert len(filas)==27 and sum(r[2]=='Inactivo' for r in filas)==1
    assert 'equipo' not in {r[0] for r in filas} and 'ajeno' not in {r[0] for r in filas}
    wb.close()


def test_textos_se_exportan_como_texto_sin_ejecutar_formulas(personal):
    c,*_=personal;e=nuevo(personal,nombre='=1+1');e.observacion='@SUM(1,2)';db.session.commit()
    wb,datos=libro(c)
    assert wb['Personal']['B2'].value=='=1+1' and wb['Personal']['B2'].data_type=='s'
    assert wb['Personal']['U2'].value=='@SUM(1,2)' and wb['Personal']['U2'].data_type=='s'
    with ZipFile(BytesIO(datos)) as z:assert b'<f>' not in b''.join(z.read(p) for p in z.namelist() if p.startswith('xl/worksheets/'))
    wb.close()


def test_exportacion_admin_y_organizacion_activa(personal):
    c,_,_,_,_,mem=personal;nuevo(personal);mem.rol='carga'
    r=c.get('/admin/personal/exportar');assert r.status_code==302 and r.mimetype!='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def test_exportacion_vacia_y_fuentes_no_preparadas_se_informan_sin_crear_tablas(personal):
    from sqlalchemy import inspect
    c,*_=personal
    db.metadata.drop_all(db.engine,tables=[PersonalAcceso.__table__])
    wb,_=libro(c)
    assert wb['Personal'].max_row==1 and wb['Horarios'].max_row==1
    assert 'Accesos' in list(wb['Información'].values)[-1][1]
    assert not inspect(db.engine).has_table(PersonalAcceso.__tablename__)
    wb.close()
