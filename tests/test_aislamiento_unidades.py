"""CRM y legajos compartidos: separación real en DB, rutas, planes y exportaciones."""
from types import SimpleNamespace
from io import BytesIO
import pytest
from extensions import db
from test_producto_ficha_maestra import entorno
from test_personal_admin import personal,nuevo
from models.unidad_negocio import UnidadNegocio
from models.usuario_sistema import UsuarioSistema
from models.cliente_crm import ClienteCRM
from models.cliente_unidad import ClienteUnidad
from models.cliente_identidad_canal import ClienteIdentidadCanal
from models.etapa_crm import EtapaCRM
from models.oportunidad_crm import OportunidadCRM
from models.actividad_crm import ActividadCRM
from models.lote_importacion_crm import LoteImportacionCRM
from models.modulo_organizacion import ModuloOrganizacion
from models.fuentes_costo_productivo import EmpleadoProductivo
from services.clientes_unidad import vincular,procesar
from scripts.preparar_clientes_unidad import preparar


@pytest.fixture()
def crm(entorno,monkeypatch):
    import modules.admin.crm.routes as rutas
    app,org,otra,a=entorno
    a.activa=True;b=UnidadNegocio(organizacion_id=org.id,codigo='b',nombre='Unidad B',activa=True)
    db.session.add(b);db.session.commit()
    tablas=[ClienteCRM.__table__,ClienteIdentidadCanal.__table__,EtapaCRM.__table__,OportunidadCRM.__table__,ActividadCRM.__table__,LoteImportacionCRM.__table__,ModuloOrganizacion.__table__]
    db.metadata.create_all(db.engine,tables=tablas)
    clientes=[ClienteCRM(organizacion_id=org.id,unidad_negocio_id=u,codigo=cod,nombre=nom,activo=False) for u,cod,nom in [(a.id,'a','Cliente A'),(b.id,'b','Cliente B'),(None,'g','General sin asignar')]]
    db.session.add_all(clientes);db.session.commit()
    preparar(db.engine);preparar(db.engine)
    mem=SimpleNamespace(rol='admin',organizacion_id=org.id,organizacion=org)
    monkeypatch.setattr(rutas,'resolver_tenant_usuario',lambda *args,**kw:mem)
    modelos={M.__name__:M for M in [UnidadNegocio,ClienteCRM,ClienteIdentidadCanal,EtapaCRM,OportunidadCRM,ActividadCRM,LoteImportacionCRM,ModuloOrganizacion]}
    app.register_blueprint(rutas.crear_blueprint_crm(dependencias=dict(db=db,modelos=modelos,usuario_actual=lambda:SimpleNamespace(id=None,username='admin'),UsuarioOrganizacion=object,login_required=lambda f:f,registrar_auditoria=lambda *a,**k:None)))
    app.add_url_rule('/',endpoint='inicio',view_func=lambda:'Inicio')
    app.url_build_error_handlers.append(lambda e,endpoint,values:'/sin-uso/'+endpoint)
    app.context_processor(lambda:dict(rol_actual=mem.rol,csrf_token_value='test',csp_nonce='test',usuario_logueado=SimpleNamespace(nombre='Admin',id=1)))
    app.jinja_env.globals.update(puede_crear_pedido=lambda:False,puede_ver_historico=lambda:False,alertas_operativas=lambda:[])
    c=app.test_client();c.environ_base['HTTP_X_FIERRO_UNIDAD']=str(a.id)
    with c.session_transaction() as s:s['unidad_negocio_id']=a.id
    yield c,org,otra,a,b,clientes,modelos,mem
    db.session.rollback();db.metadata.drop_all(db.engine,tables=[ClienteUnidad.__table__]+list(reversed(tablas)))


def test_crm_lista_y_estado_solo_unidad(crm):
    c,org,_,a,b,cs,m,*_=crm
    r=c.get('/admin/crm?vista=clientes');assert r.status_code==200,r.text
    assert 'Cliente A' in r.text and 'Cliente B' not in r.text and 'General sin asignar' not in r.text
    r=c.post('/admin/crm/guardar',data=dict(accion='toggle_cliente',cliente_id=cs[1].id))
    assert 'error=' in r.location and not cs[1].activo
    r=c.post('/admin/crm/guardar',data=dict(accion='vincular_cliente',codigo_existente='b'))
    assert 'ok=' in r.location
    r=c.post('/admin/crm/guardar',data=dict(accion='toggle_cliente',cliente_id=cs[1].id))
    assert 'ok=' in r.location
    assert ClienteUnidad.query.filter_by(cliente_id=cs[1].id,unidad_negocio_id=a.id).one().activo
    assert not ClienteUnidad.query.filter_by(cliente_id=cs[1].id,unidad_negocio_id=b.id).one().activo
    assert ClienteCRM.query.count()==3


def test_crm_documentos_compartidos_y_formulario_obsoleto(crm):
    c,org,_,a,b,cs,m,*_=crm
    vincular(cs[0],org.id,b.id,db.session)
    act=ActividadCRM(organizacion_id=org.id,unidad_negocio_id=b.id,cliente_crm_id=cs[0].id,asunto='Nota privada B',tipo='nota',estado='pendiente')
    db.session.add(act);db.session.commit()
    assert 'Nota privada B' not in c.get('/admin/crm').text
    r=c.post('/admin/crm/guardar',data=dict(accion='estado_actividad',actividad_id=act.id,estado='completada'))
    assert 'error=' in r.location and act.estado=='pendiente'
    r=c.post('/admin/crm/guardar',data=dict(accion='crear_cliente',codigo='obsoleto',nombre='No crear',unidad_contexto_id=b.id))
    assert r.status_code==302 and ClienteCRM.query.filter_by(codigo='obsoleto').count()==0


def test_crm_alta_actividad_e_importacion_con_unidad(crm):
    from services.importacion_offline_crm import previsualizar_importacion
    from services.confirmacion_importacion_crm import confirmar_importacion
    c,org,_,a,b,cs,m,*_=crm
    r=c.post('/admin/crm/guardar',data=dict(accion='crear_cliente',codigo='nuevo',nombre='Nuevo',tipo='persona'))
    assert 'ok=' in r.location,r.location
    nuevo_c=ClienteCRM.query.filter_by(codigo='nuevo').one()
    assert nuevo_c.unidad_negocio_id==a.id and ClienteUnidad.query.filter_by(cliente_id=nuevo_c.id,unidad_negocio_id=a.id).count()==1
    r=c.post('/admin/crm/guardar',data=dict(accion='crear_actividad',cliente_id=nuevo_c.id,tipo='nota',asunto='Nota A'))
    assert 'ok=' in r.location,r.location
    assert ActividadCRM.query.filter_by(asunto='Nota A').one().unidad_negocio_id==a.id
    plan=previsualizar_importacion(b'tipo,codigo,nombre\ncliente,importado,Importado\n',organizacion_id=org.id,unidades=[a],etapas=[],clientes=[cs[0]],identidades=[],unidad_negocio_id=a.id)
    assert plan['unidad_negocio_id']==a.id and plan['filas'][0]['unidad_negocio_id']==a.id
    with pytest.raises(ValueError,match='unidad activa'):
        confirmar_importacion(plan,organizacion_id=org.id,usuario=SimpleNamespace(id=None),nombre_archivo='x',clientes_existentes=[],modelos=m,db_session=db.session,unidad_negocio_id=b.id)


def test_crm_evidencia_de_lote_ajeno_no_se_descarga(crm):
    c,org,_,a,b,cs,m,*_=crm
    lote=LoteImportacionCRM(organizacion_id=org.id,unidad_negocio_id=b.id,nombre_archivo='Privado B',huella_documento='x'*64,huella_plan='y'*64,evidencia_json='{}')
    db.session.add(lote);db.session.commit()
    assert c.get(f'/admin/crm/lotes/{lote.id}/evidencia').status_code==404
    assert 'Privado B' not in c.get('/admin/crm/importacion-offline').text


def test_personal_privado_compartido_y_exportacion(personal):
    from openpyxl import load_workbook
    from services.personal_unidad import vincular
    from services.fuentes_costo_productivo import crear_empleado
    c,org,_,a,m,*_=personal
    b=UnidadNegocio(organizacion_id=org.id,codigo='b',nombre='Náutica B',activa=True);db.session.add(b);db.session.commit()
    propio=nuevo(personal)
    ajeno=crear_empleado(organizacion_id=org.id,unidad_negocio_id=b.id,codigo='solo-b',nombre='Sólo unidad B',sector='B',puesto='Operario',observacion='',Organizacion=m['Organizacion'],UnidadNegocio=UnidadNegocio,EmpleadoProductivo=EmpleadoProductivo,db_session=db.session)
    assert 'Sólo unidad B' not in c.get('/admin/personal').text
    assert c.get(f'/admin/personal/{ajeno.id}').status_code==404
    vincular(ajeno.codigo,org.id,a.id,m,db.session);db.session.commit()
    assert c.get(f'/admin/personal/{ajeno.id}').status_code==200
    r=c.get('/admin/personal/exportar');assert r.status_code==200
    wb=load_workbook(BytesIO(r.data));assert wb['Personal'].max_row==3
    assert 'Unidad activa' in str(list(wb['Información'].values));wb.close()
    with c.session_transaction() as s:s['unidad_negocio_id']=b.id
    r=c.post('/admin/personal',data=dict(_csrf_token='csrf-personal',accion='crear',codigo='stale',nombre='Stale',sector='B',activo='1',unidad_contexto_id=a.id))
    assert r.status_code==302 and EmpleadoProductivo.query.filter_by(codigo='stale').count()==0
    assert c.get(f'/admin/personal/{propio.id}').status_code==404


def test_gastos_repartidos_no_duplican_deuda_y_rechazan_operacion_ajena(entorno):
    from datetime import date
    from models.fuentes_costo_productivo import CostoFijoProductivo,CostoFijoDistribucionVersion
    from services.gastos_unidad import consulta,validar_accion
    from services.origenes_proyectados_tesoreria import consolidar_origenes
    app,org,_,a=entorno
    tablas=[CostoFijoProductivo.__table__,CostoFijoDistribucionVersion.__table__]
    db.metadata.create_all(db.engine,tables=tablas)
    b=UnidadNegocio(organizacion_id=org.id,codigo='b',nombre='B',activa=True);db.session.add(b);db.session.flush()
    gasto=CostoFijoProductivo(organizacion_id=org.id,unidad_negocio_id=a.id,codigo='alquiler',nombre='Alquiler compartido',categoria='Infraestructura',criterio_distribucion='porcentaje')
    otro=CostoFijoProductivo(organizacion_id=org.id,unidad_negocio_id=b.id,codigo='privado-b',nombre='Privado B',categoria='Infraestructura',criterio_distribucion='sin_distribuir')
    general=CostoFijoProductivo(organizacion_id=org.id,codigo='sin-unidad',nombre='Sin responsable',categoria='Infraestructura',criterio_distribucion='sin_distribuir')
    db.session.add_all([gasto,otro,general]);db.session.flush()
    d=CostoFijoDistribucionVersion(organizacion_id=org.id,costo_fijo_id=gasto.id,unidad_negocio_id=b.id,numero_revision=1,porcentaje_asignacion=40,porcentaje_productivo=100,ubicacion_costo='Taller',vigente=True,vigente_desde=__import__('datetime').datetime(2026,10,27))
    db.session.add(d);db.session.commit()
    assert {x.id for x in consulta(CostoFijoProductivo,CostoFijoDistribucionVersion,org.id,a.id).all()}=={gasto.id}
    assert {x.id for x in consulta(CostoFijoProductivo,CostoFijoDistribucionVersion,org.id,b.id).all()}=={gasto.id,otro.id}
    with pytest.raises(ValueError,match='responsable'):
        validar_accion('crear_obligacion_costo',{'costo_fijo_id':gasto.id},org.id,b.id,{'CostoFijoProductivo':CostoFijoProductivo})
    obligacion=SimpleNamespace(id=1,organizacion_id=org.id,costo_fijo=gasto,estado='pendiente',pagos=[],importe_centavos=10000,fecha_vencimiento=date(2026,11,1))
    for uid,cantidad in [(a.id,1),(b.id,0)]:
        r=consolidar_origenes(organizacion_id=org.id,unidad_negocio_id=uid,obligaciones=[obligacion],facturas=[],ventas=[],proyecciones_existentes=[])
        assert len(r['candidatos'])==cantidad
    db.session.rollback();db.metadata.drop_all(db.engine,tables=list(reversed(tablas)))


def test_contexto_financiero_rechaza_formulario_anterior(entorno):
    from services.alcance_unidad import validar_formulario
    app,*_=entorno
    with app.test_request_context('/admin/tesoreria/guardar',method='POST',data={'unidad_contexto_id':'1'}):
        with pytest.raises(ValueError,match='unidad activa'):validar_formulario(2)
    with app.test_request_context('/admin/tesoreria/guardar',method='POST',headers={'X-Fierro-Unidad':'2'}):
        validar_formulario(2)


def test_confirmar_plan_crm_guarda_lote_y_relaciones_en_su_unidad(crm):
    from services.importacion_offline_crm import previsualizar_importacion
    from services.confirmacion_importacion_crm import confirmar_importacion
    c,org,_,a,b,cs,m,*_=crm
    plan=previsualizar_importacion(b'tipo,codigo,nombre\ncliente,importado,Importado\nactividad,importado,Nota importada\n',organizacion_id=org.id,unidades=[a],etapas=[],clientes=[cs[0]],identidades=[],unidad_negocio_id=a.id)
    assert not any(f['errores'] for f in plan['filas']),plan
    lote=confirmar_importacion(plan,organizacion_id=org.id,usuario=SimpleNamespace(id=None),nombre_archivo='archivo.csv',clientes_existentes=[cs[0]],modelos=m,db_session=db.session,unidad_negocio_id=a.id)
    assert lote.unidad_negocio_id==a.id
    cliente=ClienteCRM.query.filter_by(codigo='importado').one()
    assert ClienteUnidad.query.filter_by(cliente_id=cliente.id,unidad_negocio_id=a.id).count()==1
    assert ActividadCRM.query.filter_by(cliente_crm_id=cliente.id).one().unidad_negocio_id==a.id
    with c.session_transaction() as s:s['unidad_negocio_id']=b.id
    assert c.get(f'/admin/crm/lotes/{lote.id}/evidencia').status_code==404


def test_preparacion_clientes_amplia_esquema_legacy_sin_inventar_asignaciones():
    from sqlalchemy import create_engine,text,inspect
    engine=create_engine('sqlite://')
    with engine.begin() as c:
        for ddl in [
            'CREATE TABLE organizacion(id INTEGER PRIMARY KEY)',
            'CREATE TABLE unidad_negocio(id INTEGER PRIMARY KEY,organizacion_id INTEGER)',
            'CREATE TABLE cliente_crm(id INTEGER PRIMARY KEY,organizacion_id INTEGER,unidad_negocio_id INTEGER,activo BOOLEAN)',
            'CREATE TABLE oportunidad_crm(id INTEGER PRIMARY KEY,organizacion_id INTEGER,cliente_crm_id INTEGER,unidad_negocio_id INTEGER)',
            'CREATE TABLE actividad_crm(id INTEGER PRIMARY KEY,organizacion_id INTEGER,cliente_crm_id INTEGER,oportunidad_crm_id INTEGER)',
            'CREATE TABLE cliente_identidad_canal(id INTEGER PRIMARY KEY,organizacion_id INTEGER,cliente_crm_id INTEGER)',
            'CREATE TABLE lote_importacion_crm(id INTEGER PRIMARY KEY,organizacion_id INTEGER)']:
            c.execute(text(ddl))
        for sql in ['INSERT INTO organizacion VALUES(1)','INSERT INTO unidad_negocio VALUES(1,1),(2,1)',
            'INSERT INTO cliente_crm VALUES(1,1,1,1),(2,1,NULL,0)',
            'INSERT INTO oportunidad_crm VALUES(1,1,1,2)',
            'INSERT INTO actividad_crm VALUES(1,1,1,1),(2,1,2,NULL)',
            'INSERT INTO cliente_identidad_canal VALUES(1,1,1),(2,1,2)',
            'INSERT INTO lote_importacion_crm VALUES(1,1)']:
            c.execute(text(sql))
    preparar(engine);preparar(engine)
    with engine.connect() as c:
        assert c.execute(text('SELECT cliente_id,unidad_negocio_id FROM cliente_unidad ORDER BY unidad_negocio_id')).all()==[(1,1),(1,2)]
        assert c.execute(text('SELECT unidad_negocio_id FROM actividad_crm ORDER BY id')).all()==[(2,),(None,)]
        assert c.execute(text('SELECT unidad_negocio_id FROM cliente_identidad_canal ORDER BY id')).all()==[(1,),(None,)]
        assert c.execute(text('SELECT unidad_negocio_id FROM lote_importacion_crm')).scalar() is None
        assert c.execute(text('SELECT COUNT(*) FROM cliente_crm')).scalar()==2
    engine.dispose()


def test_ciclo_ipc_acotado_no_aplica_propuestas_de_otra_unidad(monkeypatch):
    from datetime import date
    import services.ajustes_costos_ipc as ipc
    a=SimpleNamespace(id=1,proximo_ajuste=date(2027,3,1),periodo_ipc_inicio=date(2026,9,1),periodo_ipc_final=date(2027,2,1),frecuencia_meses=6)
    pa,pb=SimpleNamespace(regla_id=1),SimpleNamespace(regla_id=2)
    class Query:
        def filter_by(self,**kw):return self
        def all(self):return [pa,pb]
    aplicadas=[];generadas=[];consultas=[]
    monkeypatch.setattr(ipc,'actualizar_indices_oficiales',lambda **kw:consultas.append(kw))
    monkeypatch.setattr(ipc,'generar_propuestas',lambda **kw:generadas.extend(kw['reglas']))
    monkeypatch.setattr(ipc,'aplicar_propuesta',lambda p,**kw:aplicadas.append(p) or 1)
    modelos=dict(ReglaAjusteIPCProductivo=object,PropuestaAjusteIPCProductivo=SimpleNamespace(query=Query()),IndiceIPCOficial=object,CostoFijoVersion=object)
    assert ipc.ejecutar_ciclo_ipc(modelos=modelos,db_session=object(),reglas=[a])==1
    assert aplicadas==[pa] and generadas==[a] and len(consultas)==1


def test_importacion_de_personal_no_acepta_legajo_general_sin_vinculo(personal):
    from services.importacion_fuentes_costeo import previsualizar_fuentes,definicion
    c,org,_,a,m,*_=personal
    general=EmpleadoProductivo(organizacion_id=org.id,codigo='sin-asignacion',nombre='Privado sin asignar',sector='Sin unidad',tipo_registro='empleado',unidad_negocio_id=None)
    db.session.add(general);db.session.commit()
    campos=list(definicion('empleados')['campos'])
    datos={k:'1' for k in campos};datos.update(codigo=general.codigo,nombre='Nombre del archivo',sector='Sector del archivo')
    vista=previsualizar_fuentes('empleados',[dict(numero=1,valores=[datos[k] for k in campos])],{str(i):k for i,k in enumerate(campos)},organizacion_id=org.id,unidad_negocio_id=a.id,modelos=m,unidad_estricta=True)
    assert vista[0]['accion']=='rechazado' and 'unidad activa' in ' '.join(vista[0]['errores'])
    assert 'Privado sin asignar' not in str(vista)


def test_agenda_general_usa_mismos_vinculos_que_personal(personal,monkeypatch):
    from flask import current_app
    import modules.admin.personas.routes as rutas
    c,org,_,a,m,mem=personal
    monkeypatch.setattr(rutas,'resolver_tenant_usuario',lambda *a,**kw:mem)
    current_app.register_blueprint(rutas.crear_blueprint_personas(dependencias=dict(db=db,modelos=m,usuario_actual=lambda:SimpleNamespace(id=1),UsuarioOrganizacion=m['UsuarioOrganizacion'],login_required=lambda f:f,registrar_auditoria=lambda *a,**k:None)))
    e=nuevo(personal)
    general=EmpleadoProductivo(organizacion_id=org.id,codigo='oculto',nombre='No debe verse en agenda',sector='Sin unidad',tipo_registro='empleado',unidad_negocio_id=None)
    db.session.add(general);db.session.commit()
    r=c.get('/admin/personas');assert r.status_code==200,r.text
    assert e.nombre in r.text and general.nombre not in r.text
