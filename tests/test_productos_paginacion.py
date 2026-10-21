import json
import re
from io import BytesIO
from openpyxl import load_workbook
from extensions import db
from models.producto import Producto
from test_producto_ficha_maestra import entorno, configurar_rutas
from test_producto_datos_generales import completo
from services.productos_consultas import obtener_panel_productos_plataforma


def poblar(completo):
    app,org,otra,*_=completo
    db.session.add_all([Producto(organizacion_id=org.id,sku=f'P-{n:03}',descripcion='Producto') for n in range(61)])
    db.session.add(Producto(organizacion_id=otra.id,sku='AJENO',descripcion='Privado'))
    db.session.commit()
    configurar_rutas(app,org)
    return app.test_client(),org


def filas(respuesta):
    assert respuesta.status_code==200
    bloque=re.search(r'id="productos-datos">(.*?)</script>',respuesta.text,re.S)
    return json.loads(bloque[1])['filas']


def test_paginas_10_25_50_y_preferencia(completo):
    c,org=poblar(completo)
    r=c.get('/admin/productos');assert len(filas(r))==10
    assert 'Mostrando 1–10 de 61' in r.text and 'Página 1 de 7' in r.text
    a=filas(c.get('/admin/productos?pagina=2'));assert a[0]['datos']['sku']=='P-010'
    for n in (25,50):
        r=c.get(f'/admin/productos?por_pagina={n}');assert len(filas(r))==n
        assert len(filas(c.get('/admin/productos')))==n
    r=c.get('/admin/productos?pagina=2');assert len(filas(r))==11
    assert 'Mostrando 51–61 de 61' in r.text
    assert 'rel="prev"' in r.text and 'rel="next"' not in r.text


def test_parametros_invalidos_y_busqueda_global(completo):
    c,org=poblar(completo)
    for pagina in ('-1','texto','0'):
        assert filas(c.get('/admin/productos?pagina='+pagina))[0]['datos']['sku']=='P-000'
    assert filas(c.get('/admin/productos?pagina=999'))[0]['datos']['sku']=='P-060'
    for tamano in ('0','100000','texto','-1'):
        assert len(filas(c.get('/admin/productos?por_pagina='+tamano)))==10
    r=c.get('/admin/productos?sku=P-05&por_pagina=25')
    assert len(filas(r))==10 and 'Mostrando 1–10 de 10' in r.text
    assert not filas(c.get('/admin/productos?sku=AJENO'))
    assert 'Mostrando 0–0 de 0' in c.get('/admin/productos?sku=NO-EXISTE').text
    r=c.get('/admin/productos?sku=P-&por_pagina=25&pagina=2')
    assert 'sku=P-' in r.text and 'por_pagina=25' in r.text
    assert len(filas(r))==25


def test_exportacion_incluye_todas_las_paginas(completo):
    c,org=poblar(completo)
    assert len(filas(c.get('/admin/productos?por_pagina=10&pagina=3')))==10
    r=c.get('/admin/productos/exportar.xlsx')
    assert r.status_code==200
    w=load_workbook(BytesIO(r.data),read_only=True)
    assert w['Productos'].max_row==62
    w.close()
