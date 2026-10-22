from io import BytesIO
import re
import pytest
from openpyxl import load_workbook
from werkzeug.datastructures import FileStorage
from extensions import db
from models.producto import Producto
from services import planilla_productos
from test_producto_ficha_maestra import entorno,crear,configurar_rutas
from test_producto_datos_generales import completo


def cliente(c):
    crear(c[1]);configurar_rutas(c[0],c[1]);return c[0].test_client()


def test_menus_y_tareas_aisladas_sin_formularios_en_listado(completo):
    c=cliente(completo);r=c.get('/admin/productos')
    assert r.status_code==200
    assert 'data-productos-menu="productos-menu-exportar"' in r.text
    assert 'data-productos-menu="productos-menu-importar"' in r.text
    assert 'data-import-open' not in r.text
    assert 'data-auto-open' not in r.text
    assert len(re.findall('data-productos-dialogo=',r.text))==7
    assert '<summary>Importación masiva</summary>' not in r.text
    assert 'admin_productos_tareas.js' in r.text


@pytest.mark.parametrize('accion,tarea,archivo',[
 ('previsualizar_selectiva','importar-selectiva','archivo_selectivo'),
 ('previsualizar_planilla','importar-editable','archivo_planilla'),
 ('previsualizar_importacion','importar-altas','archivo_productos'),
 ('previsualizar_fabricacion','importar-fabricacion','archivo_fabricacion'),
 ('previsualizar_logistica_tn','importar-tiendanube','archivo_logistica_tn')])
def test_error_reabre_la_emergente_correcta(completo,accion,tarea,archivo):
    c=cliente(completo)
    r=c.post('/admin/productos',data={'accion':accion,'_csrf_token':'csrf-test'})
    assert r.status_code==200
    marcados=re.findall(r'data-productos-dialogo="([^"]+)" data-auto-open',r.text)
    assert marcados==[tarea]
    cuerpo=r.text.split(f'data-productos-dialogo="{tarea}"',1)[1].split('</dialog>',1)[0]
    assert 'role="alert"' in cuerpo and f'name="{archivo}"' in cuerpo
    assert Producto.query.count()==1


def test_base_completa_consulta_y_editable_misma_base_y_importacion_correcta(completo):
    c=cliente(completo)
    edit=c.get('/admin/productos/exportar.xlsx')
    consulta=c.get('/admin/productos/exportar.xlsx?uso=consulta')
    assert edit.status_code==consulta.status_code==200
    a=load_workbook(BytesIO(edit.data));b=load_workbook(BytesIO(consulta.data))
    assert list(a['Productos'].values)==list(b['Productos'].values)
    assert a['Control']['A2'].value=='FIERRO-PRODUCTOS-1'
    assert b['Control']['A2'].value=='FIERRO-PRODUCTOS-CONSULTA-1'
    assert all(ws.protection.sheet for ws in b)
    assert all(cell.protection.locked for ws in b for row in ws for cell in row)
    assert not a['Productos']['D2'].protection.locked
    assert 'base_completa_productos_fierro.xlsx' in consulta.headers['Content-Disposition']
    with pytest.raises(ValueError,match='base de consulta'):
        planilla_productos.leer(FileStorage(BytesIO(consulta.data),filename='consulta.xlsx'),completo[1].id)
    assert planilla_productos.leer(FileStorage(BytesIO(edit.data),filename='editable.xlsx'),completo[1].id)


def test_exportacion_fallida_reabre_su_tarea(completo,monkeypatch):
    c=cliente(completo)
    def fallo(**kw):raise ValueError('Revisá el maestro')
    monkeypatch.setattr(planilla_productos,'exportar',fallo)
    r=c.get('/admin/productos/exportar.xlsx?uso=consulta',follow_redirects=True)
    assert r.status_code==200
    assert 'data-productos-dialogo="exportar-completa" data-auto-open' in r.text
    assert 'Revisá el maestro' in r.text
